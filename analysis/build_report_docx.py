# -*- coding: utf-8 -*-
# Builds the Word version of the sales analysis report from the report HTML.
# Usage: python build_report_docx.py <input.html> <output.docx>
import sys, re
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0B, 0x3D, 0x66); GREEN = RGBColor(0x1E, 0x7A, 0x46); RED = RGBColor(0xB3, 0x26, 0x1E); GREY = RGBColor(0x66, 0x66, 0x66)
src, out = sys.argv[1], sys.argv[2]
soup = BeautifulSoup(open(src, encoding='utf-8').read(), 'html.parser')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.9); sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(2.0)
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15

def shade(el, fill):
    pr = el.get_or_add_tcPr() if hasattr(el, 'get_or_add_tcPr') else el.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); pr.append(sh)
def border(p, side, color='0B3D66', sz=12, space=4):
    pPr = p._p.get_or_add_pPr(); b = pPr.find(qn('w:pBdr'))
    if b is None: b = OxmlElement('w:pBdr'); pPr.append(b)
    e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color); b.append(e)

def add_inline(p, node, size=None, color=None, bold=False, italic=False):
    for ch in node.children:
        if isinstance(ch, NavigableString):
            t = str(ch).replace('\n', ' ')
            t = re.sub(r'\s+', ' ', t)
            if not t: continue
            r = p.add_run(t); r.bold = bold; r.italic = italic
            if size: r.font.size = Pt(size)
            if color is not None: r.font.color.rgb = color
        else:
            b, i, c = bold, italic, color
            if ch.name in ('b', 'strong'): b = True
            if ch.name in ('i', 'em'): i = True
            cl = ch.get('class') or []
            if 'pos' in cl: c = GREEN
            if 'neg' in cl: c = RED
            if ch.name == 'br':
                p.add_run().add_break(); continue
            add_inline(p, ch, size, c, b, i)

def para(node, size=10, color=None, after=6, keep=False):
    p = doc.add_paragraph(); add_inline(p, node, size, color)
    p.paragraph_format.space_after = Pt(after)
    if keep: p.paragraph_format.keep_with_next = True
    return p

def heading(text, level):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.color.rgb = NAVY
    p.paragraph_format.keep_with_next = True
    if level == 1:
        r.font.size = Pt(19); p.paragraph_format.space_after = Pt(2)
    elif level == 2:
        r.font.size = Pt(14); p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6); border(p, 'bottom')
    else:
        r.font.size = Pt(11); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    return p

def add_table(t):
    rows = t.find_all('tr'); ncol = max(len(r.find_all(['td', 'th'])) for r in rows)
    tb = doc.add_table(rows=len(rows), cols=ncol); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    total = 17.2
    first = 5.6 if ncol >= 6 else 6.2 if ncol == 4 else 6.8
    if ncol <= 3: first = 8.0
    widths = [first] + [(total - first) / (ncol - 1)] * (ncol - 1)
    if ncol == 4 and any('Likely cause' in (c.get_text()) for c in rows[0].find_all('th')): widths = [3.6, 3.0, 1.5, 1.5, 1.5, 6.1][:ncol]
    if ncol == 6 and any('Likely cause' in (c.get_text()) for c in rows[0].find_all('th')): widths = [3.6, 3.0, 1.4, 1.4, 1.5, 6.3]
    for ci, w in enumerate(widths[:ncol]): tb.columns[ci].width = Cm(w)
    for ri, tr in enumerate(rows):
        cells = tr.find_all(['td', 'th']); trcls = tr.get('class') or []
        trPr = tb.rows[ri]._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
        if ri == 0:
            th = OxmlElement('w:tblHeader'); trPr.append(th)
        for ci in range(ncol):
            cell = tb.cell(ri, ci); cell.width = Cm(widths[ci] if ci < len(widths) else 2)
            if ci >= len(cells): continue
            c = cells[ci]; cl = c.get('class') or []
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
            hdr = c.name == 'th'
            color = RED if 'neg' in cl else GREEN if 'pos' in cl else None
            if hdr: color = RGBColor(0xFF, 0xFF, 0xFF)
            add_inline(p, c, 8.5, color, bold=hdr or 'tot' in trcls, italic='sub' in trcls)
            left = (ci == 0) or ('l' in cl)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.RIGHT
            if 'sub' in trcls and ci == 0: p.paragraph_format.left_indent = Cm(0.5)
            if ri < len(rows) - 1: p.paragraph_format.keep_with_next = True
            tcPr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side in ('top', 'bottom'):
                m = OxmlElement('w:' + side); m.set(qn('w:w'), '50'); m.set(qn('w:type'), 'dxa'); mar.append(m)
            for side in ('left', 'right'):
                m = OxmlElement('w:' + side); m.set(qn('w:w'), '80'); m.set(qn('w:type'), 'dxa'); mar.append(m)
            tcPr.append(mar)
            if hdr: shade(cell._tc, '0B3D66')
            elif ri % 2 == 0: shade(cell._tc, 'F7F9FB')
            bd = OxmlElement('w:tcBorders'); b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'E4E4E4'); bd.append(b)
            if 'tot' in trcls:
                tp = OxmlElement('w:top'); tp.set(qn('w:val'), 'single'); tp.set(qn('w:sz'), '12'); tp.set(qn('w:color'), '0B3D66'); bd.append(tp)
            tcPr.append(bd)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4); sp.paragraph_format.line_spacing = 0.6

body = soup.body
for el in body.children:
    if isinstance(el, NavigableString): continue
    n = el.name; cl = el.get('class') or []
    if n == 'h1': heading(el.get_text(), 1)
    elif n == 'h2': heading(el.get_text(), 2)
    elif n == 'h3': heading(el.get_text(), 3)
    elif n == 'p':
        if 'note' in cl: para(el, 8, GREY)
        else: para(el)
    elif n == 'div' and 'meta' in cl:
        para(el, 8, GREY, 10)
    elif n == 'div' and ('key' in cl or 'gap' in cl or 'warn' in cl):
        # split on <br><br> into paragraphs inside a shaded, left-ruled block
        parts, cur = [], []
        for ch in el.children:
            if getattr(ch, 'name', None) == 'br':
                if cur and cur[-1] == 'BR':
                    cur.pop(); parts.append(cur); cur = []
                else: cur.append('BR')
            else: cur.append(ch)
        parts.append(cur)
        for k, part in enumerate(parts):
            part = [x for x in part if x != 'BR']
            if not part: continue
            p = doc.add_paragraph(); shade(p._p, 'F7F9FB'); border(p, 'left', '0B3D66', 24, 8)
            p.paragraph_format.left_indent = Cm(0.3); p.paragraph_format.space_after = Pt(0 if k < len(parts) - 1 else 8)
            p.paragraph_format.keep_together = True
            wrapper = BeautifulSoup('<x></x>', 'html.parser').x
            for x in part: wrapper.append(x.__copy__() if hasattr(x, '__copy__') else x)
            add_inline(p, wrapper, 9.5)
    elif n == 'ul':
        for li in el.find_all('li', recursive=False):
            p = doc.add_paragraph(style='List Bullet'); add_inline(p, li, 10); p.paragraph_format.space_after = Pt(3)
    elif n == 'table': add_table(el)
    elif n == 'dl':
        for ch in el.children:
            if isinstance(ch, NavigableString): continue
            if ch.name == 'dt':
                p = doc.add_paragraph(); r = p.add_run(ch.get_text()); r.bold = True; r.font.color.rgb = NAVY
                p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
            elif ch.name == 'dd':
                para(ch, 10, None, 4)

# footer: title left, page X of Y right
def field(run, instr):
    for t, txt in (('begin', None), (None, instr), ('end', None)):
        if t:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); run._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; run._r.append(e)
fp = sec.footer.paragraphs[0]; fp.text = ''
tabs = fp.paragraph_format.tab_stops; from docx.enum.text import WD_TAB_ALIGNMENT; tabs.add_tab_stop(Cm(17.2), WD_TAB_ALIGNMENT.RIGHT)
r = fp.add_run('JPS Sales Analysis — Attribution & Requirements Review\tPage '); r.font.size = Pt(8); r.font.color.rgb = GREY
r2 = fp.add_run(); r2.font.size = Pt(8); r2.font.color.rgb = GREY; field(r2, 'PAGE')
r3 = fp.add_run(' of '); r3.font.size = Pt(8); r3.font.color.rgb = GREY
r4 = fp.add_run(); r4.font.size = Pt(8); r4.font.color.rgb = GREY; field(r4, 'NUMPAGES')
doc.core_properties.title = 'JPS Sales Analysis — Attribution & Requirements Review'
doc.save(out)
print('saved', out)
