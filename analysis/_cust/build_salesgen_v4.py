# -*- coding: utf-8 -*-
# Monthly billed sales and net generation FY2026-FY2028.
# FY2026: actuals through September, full year held at the financials LE (billed 3,299.661 GWh, net generation 4,526.950 GWh).
# FY2027/28: billed per the register / Driver Report, losses on the platform seasonal monthly path.
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

CLS = ['RT10', 'RT20', 'RT40', 'RT50', 'RT60-ST', 'RT70']
NAMES = {'RT10': 'RT10 Residential', 'RT20': 'RT20 Small commercial', 'RT40': 'RT40 Commercial', 'RT50': 'RT50 Large commercial',
         'RT60-ST': 'RT60 Street lighting', 'RT70': 'RT70 Industrial'}
YEARS = (2026, 2027, 2028)
MON = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
G = 1e6

# ---- billed sales: actuals Jan-Sep 2026 (platform), register for the rest ----
E = json.load(open('engine_out.json'))
live26 = {rc: [v / G for v in E['2026'][rc]] for rc in CLS}          # Jan-Sep are actual customer-billed sales
_src = open(r'C:\Projects\Sales_Platform\analysis\_build_wb_v2.py', encoding='utf-8').read()
_ns = {}
for _name in ('m26', 'm27', 'm28'):
    _i = _src.index(_name + ' = {'); _j = _src.index(chr(10) + '}' + chr(10), _i) + 3
    exec(_src[_i:_j], _ns)
reg = {y: {rc: [v / G for v in _ns['m%d' % (y - 2000)][rc]] for rc in CLS} for y in YEARS}

FIN = r'C:\Users\jwilson\Downloads\Financial Summary - August 2026  LE v2.xlsx'
_wb = load_workbook(FIN, data_only=True, read_only=True)
_pp = list(_wb['Presentation P&L'].iter_rows(values_only=True))
FY_BILLED = _pp[95][17] / 1000          # 3,299.661 GWh, FY2026 LE
FY_NG = _pp[97][17] / 1000              # 4,526.950 GWh, FY2026 LE
_pl = list(_wb['P&L 2025 - 2027'].iter_rows(values_only=True))
NG_FIN = [v / 1000 for v in _pl[65][40:52]]    # monthly net generation 2026: Jan-Aug actual, Sep-Dec plan shape
assert abs(sum(NG_FIN[:8]) * 1000 - 2952583.6) < 1, sum(NG_FIN[:8])

reg_fy = sum(sum(reg[2026][rc]) for rc in CLS)
bill = {2027: reg[2027], 2028: reg[2028]}
bill[2026] = {}
for rc in CLS:
    ytd = sum(live26[rc][:9])
    fy_cls = sum(reg[2026][rc]) * FY_BILLED / reg_fy              # class full year held at the register split of the financials total
    q4 = fy_cls - ytd
    sh = reg[2026][rc][9:]; ssum = sum(sh)
    bill[2026][rc] = live26[rc][:9] + [q4 * s / ssum for s in sh]
tot = {y: [sum(bill[y][rc][m] for rc in CLS) for m in range(12)] for y in YEARS}
assert abs(sum(tot[2026]) - FY_BILLED) < 1e-6

# ---- net generation 2026 ----
SEP_NG = 407.689                                            # estimate: read off the System Operations monthly chart, calibrated to finance net generation (+/-1.2 GWh)
ng26 = NG_FIN[:8] + [SEP_NG]
q4_ng = FY_NG - sum(ng26)
shape = NG_FIN[9:12]
ng26 += [q4_ng * s / sum(shape) for s in shape]
assert abs(sum(ng26) - FY_NG) < 1e-6
RATE = {2027: [27.75, 22.78, 29.34, 27.39, 30.72, 26.53, 27.56, 28.21, 25.35, 25.45, 24.75, 28.34],
        2028: [27.77, 22.81, 29.33, 27.38, 30.71, 26.52, 27.56, 28.21, 25.35, 25.45, 24.75, 28.34]}

NAVY, BLUE, CYAN, GREY, LIGHT = '062552', '0E4E95', '009FDA', '66788A', 'D9E1F2'
def F(**k): return Font(name='Arial', size=k.pop('size', 10), **k)
def fill(c): return PatternFill('solid', fgColor=c)

wb = Workbook()
# ---------------- billed by class ----------------
C = wb.active; C.title = 'Billed by rate class'
C['A1'] = 'Billed sales by rate class, FY2026 to FY2028 (GWh)'; C['A1'].font = F(bold=True, size=13, color=NAVY)
C['A2'] = 'January to September 2026 actual. October to December 2026 is set so FY2026 equals the financials LE. 2027 and 2028 per the register.'; C['A2'].font = F(color=GREY, italic=True)
hdr = ['Rate class']
for y in YEARS: hdr += ['%s %s' % (m, str(y)[2:]) for m in MON] + ['FY%d' % y]
for j, h in enumerate(hdr):
    x = C.cell(4, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.alignment = Alignment(horizontal='center'); x.fill = fill(NAVY)
for yi, y in enumerate(YEARS):
    base = 2 + yi * 13
    for m in range(12):
        if (y == 2026 and m >= 9) or y > 2026: C.cell(4, base + m).fill = fill(CYAN)
    C.cell(4, base + 12).fill = fill(BLUE)
for i, rc in enumerate(CLS):
    r = 5 + i
    C.cell(r, 1, NAMES[rc]).font = F(bold=True, color=NAVY)
    for yi, y in enumerate(YEARS):
        base = 2 + yi * 13
        for m in range(12):
            x = C.cell(r, base + m, bill[y][rc][m]); x.number_format = '#,##0.0'; x.font = F(color='0000FF')
        x = C.cell(r, base + 12, '=SUM(%s%d:%s%d)' % (L(base), r, L(base + 11), r)); x.number_format = '#,##0.0'; x.font = F(bold=True)
rt = 5 + len(CLS)
C.cell(rt, 1, 'Total billed sales').font = F(bold=True, color='FFFFFF')
for j in range(1, len(hdr) + 1):
    c_ = C.cell(rt, j); c_.fill = fill(NAVY)
    if j > 1:
        c_.value = '=SUM(%s5:%s%d)' % (L(j), L(j), rt - 1); c_.number_format = '#,##0.0'; c_.font = F(bold=True, color='FFFFFF')
C.column_dimensions['A'].width = 24
for j in range(2, len(hdr) + 1): C.column_dimensions[L(j)].width = 9.5
for yi in range(3): C.column_dimensions[L(2 + yi * 13 + 12)].width = 11
C.freeze_panes = 'B5'

# ---------------- monthly billed and net gen ----------------
M = wb.create_sheet('Billed and Net Gen', 0)
M['A1'] = 'Billed sales and net generation, monthly FY2026 to FY2028'; M['A1'].font = F(bold=True, size=13, color=NAVY)
M['A2'] = 'GWh. Losses = net generation less billed sales; loss % is losses as a share of net generation. Blue cells are inputs; the rest is calculated.'; M['A2'].font = F(color=GREY, italic=True)
for j, h in enumerate(['Period', 'Basis', 'Billed sales (GWh)', 'Loss % of net gen', 'Losses (GWh)', 'Net generation (GWh)']):
    x = M.cell(4, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = fill(NAVY); x.alignment = Alignment(horizontal='center', wrap_text=True)
M.row_dimensions[4].height = 30
r = 5; fy_rows = {}
for yi, y in enumerate(YEARS):
    first = r
    for m in range(12):
        col_c = L(2 + yi * 13 + m)
        M.cell(r, 1, '%s %d' % (MON[m], y)).font = F(bold=True, color=NAVY)
        x = M.cell(r, 3, "='Billed by rate class'!%s%d" % (col_c, rt)); x.number_format = '#,##0.0'; x.font = F()
        if y == 2026:
            basis = 'Actual' if m < 8 else ('Billed actual, net gen estimate' if m == 8 else 'Forecast')
            M.cell(r, 2, basis).font = F()
            x = M.cell(r, 6, ng26[m]); x.number_format = '#,##0.0'; x.font = F(color='0000FF')
            x = M.cell(r, 5, '=F%d-C%d' % (r, r)); x.number_format = '#,##0.0'; x.font = F()
            x = M.cell(r, 4, '=E%d/F%d' % (r, r)); x.number_format = '0.0%'; x.font = F()
        else:
            M.cell(r, 2, 'Forecast').font = F()
            x = M.cell(r, 4, RATE[y][m] / 100); x.number_format = '0.0%'; x.font = F(color='0000FF')
            x = M.cell(r, 5, '=C%d*D%d/(1-D%d)' % (r, r, r)); x.number_format = '#,##0.0'; x.font = F()
            x = M.cell(r, 6, '=C%d+E%d' % (r, r)); x.number_format = '#,##0.0'; x.font = F()
        if (y == 2026 and m >= 8) or y > 2026:
            for j in range(1, 7): M.cell(r, j).fill = fill('F3F6FB')
        r += 1
    last = r - 1
    M.cell(r, 1, 'FY%d' % y).font = F(bold=True, color='FFFFFF')
    for j, f in ((3, '=SUM(C%d:C%d)' % (first, last)), (5, '=SUM(E%d:E%d)' % (first, last)), (6, '=SUM(F%d:F%d)' % (first, last)), (4, '=E%d/F%d' % (r, r))):
        x = M.cell(r, j, f); x.number_format = '0.0%' if j == 4 else '#,##0.0'; x.font = F(bold=True, color='FFFFFF')
    for j in range(1, 7): M.cell(r, j).fill = fill(NAVY)
    fy_rows[y] = r
    r += 1
    if y == 2026:
        M.cell(r, 1, 'Per financials LE').font = F(italic=True, color=GREY)
        x = M.cell(r, 3, FY_BILLED); x.number_format = '#,##0.0'; x.font = F(color='0000FF')
        x = M.cell(r, 6, FY_NG); x.number_format = '#,##0.0'; x.font = F(color='0000FF')
        x = M.cell(r, 5, '=F%d-C%d' % (r, r)); x.number_format = '#,##0.0'; x.font = F()
        x = M.cell(r, 4, '=E%d/F%d' % (r, r)); x.number_format = '0.0%'; x.font = F()
        r += 1
        M.cell(r, 1, 'Difference').font = F(italic=True, color=GREY)
        for j in (3, 5, 6):
            x = M.cell(r, j, '=%s%d-%s%d' % (L(j), fy_rows[2026], L(j), r - 1)); x.number_format = '0.0;-0.0;0.0'; x.font = F(italic=True, color=GREY)
        r += 1
    r += 1
M.cell(r, 1, 'Year on year').font = F(bold=True, color=NAVY); r += 1
for j, h in ((3, 'Billed sales'), (6, 'Net generation')):
    x = M.cell(r, j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = fill(BLUE)
M.cell(r, 1).fill = fill(BLUE); r += 1
for y in (2027, 2028):
    M.cell(r, 1, 'FY%d vs FY%d' % (y, y - 1)).font = F(bold=True, color=NAVY)
    x = M.cell(r, 3, '=C%d/C%d-1' % (fy_rows[y], fy_rows[y - 1])); x.number_format = '+0.0%;-0.0%'; x.font = F()
    x = M.cell(r, 6, '=F%d/F%d-1' % (fy_rows[y], fy_rows[y - 1])); x.number_format = '+0.0%;-0.0%'; x.font = F()
    r += 1
M.column_dimensions['A'].width = 18; M.column_dimensions['B'].width = 32
for k in 'CDEF': M.column_dimensions[k].width = 16
M.freeze_panes = 'C5'

# ---------------- notes ----------------
N = wb.create_sheet('Notes')
notes = [
 ('FY2026', True),
 ('Billed sales are actual customer-billed volumes for January to September. October to December are set so the full year equals the financials LE of %s GWh, using the register split by rate class and month.' % format(FY_BILLED, ',.1f'), False),
 ('Net generation for January to August is the actual in the financials. September (%s GWh) is an estimate read from the System Operations monthly chart and calibrated to the financials net generation, accurate to about 1 GWh; replace it with the actual when issued. October to December are set so the full year equals the financials LE of %s GWh, on the financials plan shape.' % (format(SEP_NG, '.1f'), format(FY_NG, ',.1f')), False),
 ('FY2026 loss on billed sales and net generation is %.2f%%.' % ((1 - FY_BILLED / FY_NG) * 100), False),
 ('', False),
 ('FY2027 and FY2028', True),
 ('Billed sales are as in the register and Driver Report. Losses follow the seasonal monthly path in the platform assumptions, closing each year at about 27.1%.', False),
]
for i, (t, b) in enumerate(notes, 1):
    N.cell(i, 1, t).font = F(bold=b, color=NAVY if b else '000000'); N.cell(i, 1).alignment = Alignment(wrap_text=True, vertical='top')
N.column_dimensions['A'].width = 140
out = r'C:\Projects\Sales_Platform\JPS_Billed_Sales_and_Net_Gen_FY2026-28_v4.xlsx'
wb.save(out)
print('saved', out)
print('FY26 billed', round(sum(tot[2026]), 3), 'NG', round(sum(ng26), 3), 'loss', round((1 - FY_BILLED / FY_NG) * 100, 3))
print('billed by month 26', [round(v, 1) for v in tot[2026]])
print('NG by month 26', [round(v, 1) for v in ng26])
print('loss % 26', [round((1 - tot[2026][m] / ng26[m]) * 100, 1) for m in range(12)])
for y in (2027, 2028):
    Ls = [tot[y][m] * RATE[y][m] / 100 / (1 - RATE[y][m] / 100) for m in range(12)]
    print(y, round(sum(tot[y]), 1), 'NG', round(sum(tot[y]) + sum(Ls), 1), 'loss', round(sum(Ls) / (sum(tot[y]) + sum(Ls)) * 100, 3))
for rc in CLS: print(rc, round(sum(bill[2026][rc]), 1), 'Q4', [round(v, 1) for v in bill[2026][rc][9:]], 'reg Q4', [round(v, 1) for v in reg[2026][rc][9:]])
