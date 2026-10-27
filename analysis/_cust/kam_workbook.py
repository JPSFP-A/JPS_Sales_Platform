# -*- coding: utf-8 -*-
# Account workbook for the KAM forecast round (October 2026).
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from kam_model import *
from openpyxl.utils import get_column_letter as L   # after the star import: kam_model reuses the name L

NAVY, BLUE, CYAN, GREY, LIGHT, INPUT = '062552', '0E4E95', '009FDA', '66788A', 'D9E1F2', 'FFF2CC'
def F(**k): return Font(name='Arial', size=k.pop('size', 10), **k)
def fill(c): return PatternFill('solid', fgColor=c)
wb = Workbook()

def header(ws, row, heads, widths=None, color=NAVY):
    for j, h in enumerate(heads, 1):
        x = ws.cell(row, j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = fill(color); x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w

def title(ws, t, sub=None):
    ws['A1'] = t; ws['A1'].font = F(bold=True, size=13, color=NAVY)
    if sub: ws['A2'] = sub; ws['A2'].font = F(color=GREY, italic=True)

# ---------------- read me ----------------
R = wb.active; R.title = 'Read me'
title(R, 'Key account forecast round, October 2026', 'Prepared by Sales Forecasting and Analysis for the Key Account Managers')
lines = [
 ('What this workbook is', True),
 ('It shows the forecast we hold for your accounts, how last round\'s forecast compared with September actuals, and every adjustment on file for your accounts.', False),
 ('What we need from you', True),
 ('1. Adjustments sheet: for each adjustment set Keep, Change or Remove in the yellow column. If you change it, enter the new percentage or kWh and the months it applies to.', False),
 ('2. Permanent or temporary sheet: defection, expansion, renewal and win-back entries now carry forward until the next entry for the account (enter 0% to end one). Operational disruption entries apply only to the months entered. Tell us which are permanent and when any should end.', False),
 ('3. New customers sheet: confirm the start date and monthly volume, or tell us the customer is not coming.', False),
 ('4. Portfolio forecast sheet: add anything we are missing for your customers (new connections, closures, expansions, solar, self-generation).', False),
 ('', False),
 ('Sheets', True),
 ('Accuracy by KAM: September forecast against actual for each portfolio.', False),
 ('Sept account misses: every account that missed by more than 0.05 GWh.', False),
 ('Portfolio forecast: each customer, FY2025 actual and FY2026 to FY2028 forecast.', False),
 ('Adjustments: all 198 adjustments on file, with the effect on the forecast.', False),
 ('Permanent or temporary: how each customer\'s entries apply and what they do to the forecast.', False),
 ('New customers: the pipeline customers in the forecast.', False),
 ('Account forecast: monthly kWh by account, October 2026 to December 2028.', False),
 ('', False),
 ('How to read the numbers', True),
 ('Forecast last round is the account forecast as it stood after the August close, with the adjustments then on file. September actual is billed sales.', False),
 ('Net accuracy = 1 less the absolute gap between actual and forecast, divided by actual. Over and under misses cancel.', False),
 ('Account-level accuracy = 1 less the sum of each account\'s absolute miss, divided by total actual. Misses do not cancel.', False),
 ('FY2026 is January to September actual plus forecast. Portfolio totals cover assigned commercial accounts only and do not add to company sales.', False),
]
for i, (t, b) in enumerate(lines, 4):
    R.cell(i, 1, t).font = F(bold=b, color=NAVY if b else '000000'); R.cell(i, 1).alignment = Alignment(wrap_text=True, vertical='top')
R.column_dimensions['A'].width = 150

# ---------------- accuracy by KAM ----------------
S = wb.create_sheet('Accuracy by KAM')
title(S, 'September 2026: forecast last round against actual, by KAM portfolio (GWh)', 'Forecast is the account forecast after the August close. Blue cells are inputs.')
header(S, 4, ['KAM', 'Accounts', 'Forecast', 'Actual', 'Variance', 'Net accuracy', 'Sum of account misses', 'Account-level accuracy', 'Bias (variance / actual)'], [22, 10, 12, 12, 12, 13, 16, 16, 16])
r0 = 5
for i, k in enumerate(KAMS_SORTED):
    P = PORT[k]; r = r0 + i
    S.cell(r, 1, k).font = F(bold=True, color=NAVY); S.cell(r, 2, P['accounts'])
    for j, v in ((3, P['sep_fc']), (4, P['sep_act']), (7, P['sep_abs'])):
        x = S.cell(r, j, v); x.number_format = '#,##0.00'; x.font = F(color='0000FF')
    S.cell(r, 5, '=D%d-C%d' % (r, r)); S.cell(r, 6, '=1-ABS(E%d)/D%d' % (r, r)); S.cell(r, 8, '=1-G%d/D%d' % (r, r)); S.cell(r, 9, '=E%d/D%d' % (r, r))
rt = r0 + len(KAMS_SORTED)
S.cell(rt, 1, 'Total assigned').font = F(bold=True, color='FFFFFF')
S.cell(rt, 2, '=SUM(B%d:B%d)' % (r0, rt - 1))
for j in (3, 4, 7): S.cell(rt, j, '=SUM(%s%d:%s%d)' % (L(j), r0, L(j), rt - 1))
S.cell(rt, 5, '=D%d-C%d' % (rt, rt)); S.cell(rt, 6, '=1-ABS(E%d)/D%d' % (rt, rt)); S.cell(rt, 8, '=1-G%d/D%d' % (rt, rt)); S.cell(rt, 9, '=E%d/D%d' % (rt, rt))
for rr in range(r0, rt + 1):
    for j in range(2, 10):
        x = S.cell(rr, j)
        if rr == rt: x.font = F(bold=True, color='FFFFFF'); x.fill = fill(NAVY)
        elif x.font.color is None or x.font.color.rgb != '000000FF': x.font = F()
        x.number_format = '0.0%' if j in (6, 8, 9) else ('#,##0' if j == 2 else '#,##0.00;-#,##0.00')
    if rr == rt: S.cell(rr, 1).fill = fill(NAVY)
S.freeze_panes = 'B5'

# ---------------- misses ----------------
M = wb.create_sheet('Sept account misses')
title(M, 'September 2026: accounts that missed by more than 0.05 GWh', 'Sorted by size of miss. Add the reason in the last column.')
header(M, 4, ['KAM', 'Account', 'Customer', 'Rate class', 'Forecast (MWh)', 'Actual (MWh)', 'Variance (MWh)', 'Miss as % of actual', 'Reason from the KAM'], [20, 17, 36, 10, 14, 14, 14, 14, 50])
rows = [a for a in MISSES if abs(a['sep_act'] - a['sep_fc']) > 0.05]
for i, a in enumerate(rows, 5):
    M.cell(i, 1, a['kam'] or 'Not assigned'); M.cell(i, 2, a['ac']); M.cell(i, 3, a['name']); M.cell(i, 4, a['rc'])
    for j, v in ((5, a['sep_fc'] * 1000), (6, a['sep_act'] * 1000)):
        x = M.cell(i, j, v); x.number_format = '#,##0'; x.font = F(color='0000FF')
    M.cell(i, 7, '=F%d-E%d' % (i, i)).number_format = '#,##0;-#,##0'
    M.cell(i, 8, '=IF(F%d=0,"",G%d/F%d)' % (i, i, i)).number_format = '0%'
    M.cell(i, 9).fill = fill(INPUT)
    for j in (1, 2, 3, 4, 7, 8): M.cell(i, j).font = F()
M.freeze_panes = 'D5'

# ---------------- portfolio forecast ----------------
P_ = wb.create_sheet('Portfolio forecast')
title(P_, 'Forecast by customer, GWh', 'FY2026 is January to September actual plus forecast. Last round is the forecast after the August close. Add anything missing in the yellow column.')
header(P_, 4, ['KAM', 'Customer', 'Industry', 'Rate classes', 'Accounts', 'FY2025 actual', 'FY2026', 'FY2027', 'FY2028', 'FY2027 growth', 'FY2027 last round', 'Change since last round', 'Upcoming adjustments', 'KAM comment or missing information'],
       [20, 38, 32, 14, 9, 12, 11, 11, 11, 12, 13, 14, 13, 55])
upc = collections.defaultdict(int)
for g in UPCOMING: upc[(g['kam'], norm(g['name']) or g['ac'])] += 1
i = 5
for k in KAMS_SORTED:
    for c in sorted([c for c in CUST.values() if c['kam'] == k], key=lambda c: -c['live'][1]):
        P_.cell(i, 1, k); P_.cell(i, 2, c['name']); P_.cell(i, 3, c['ind']); P_.cell(i, 4, ', '.join(sorted(c['rcs']))); P_.cell(i, 5, len(c['accounts']))
        for j, v in ((6, c['fy25']), (7, c['live'][0]), (8, c['live'][1]), (9, c['live'][2]), (11, c['prev'][1])):
            x = P_.cell(i, j, v); x.number_format = '#,##0.00'; x.font = F(color='0000FF')
        P_.cell(i, 10, '=IF(G%d=0,"",H%d/G%d-1)' % (i, i, i)).number_format = '+0.0%;-0.0%'
        P_.cell(i, 12, '=H%d-K%d' % (i, i)).number_format = '+#,##0.00;-#,##0.00;0.00'
        P_.cell(i, 13, upc.get((k, norm(c['name']) or c['accounts'][0]), 0))
        P_.cell(i, 14).fill = fill(INPUT)
        for j in (1, 2, 3, 4, 5, 10, 12, 13): P_.cell(i, j).font = F()
        i += 1
P_.freeze_panes = 'C5'; P_.auto_filter.ref = 'A4:N%d' % (i - 1)

# ---------------- adjustments ----------------
A_ = wb.create_sheet('Adjustments')
title(A_, 'Adjustments on file (198 entries)', 'Effect is the whole account, so rows for the same account share it and should not be added. Choose Keep, Change or Remove in the yellow column.')
heads = ['KAM', 'Customer', 'Account', 'Month', 'Adjustment %', 'Manual kWh', 'Reason', 'How it applies', 'Justification', 'Entered by', 'Date entered', 'Status', 'Account effect FY2027 (GWh)', 'Account effect FY2028 (GWh)', 'KAM response', 'New % or kWh', 'Months it applies to', 'Comment']
header(A_, 4, heads, [20, 34, 17, 10, 12, 12, 20, 18, 60, 18, 12, 18, 14, 14, 14, 13, 20, 40])
order = sorted(adj, key=lambda a: (a['kam'] or 'zzz', a['name'] or '', a['ac'], a['year'], a['month']))
for i, a in enumerate(order, 5):
    t = (a['year'] - 2026) * 12 + a['month']
    A_.cell(i, 1, a['kam'] or 'Not assigned (no effect)'); A_.cell(i, 2, a['name'] or ''); A_.cell(i, 3, a['ac']); A_.cell(i, 4, t_label(t))
    A_.cell(i, 5, a['pct'] / 100).number_format = '+0%;-0%;0%'; A_.cell(i, 6, a['kwh']).number_format = '#,##0'
    A_.cell(i, 7, a['reason'] or 'Not stated')
    A_.cell(i, 8, 'No value' if not (a['pct'] or a['kwh']) else ('Rolls forward' if a['reason'] in ROLLS else 'Months entered only'))
    A_.cell(i, 9, a['just']); A_.cell(i, 10, a['by_name'] or 'Sales Forecasting'); A_.cell(i, 11, a['created'])
    A_.cell(i, 12, 'Upcoming' if t >= 10 else 'Already in actuals')
    if a['ac'] in effect:
        A_.cell(i, 13, effect[a['ac']][1]).number_format = '+#,##0.00;-#,##0.00;0.00'; A_.cell(i, 14, effect[a['ac']][2]).number_format = '+#,##0.00;-#,##0.00;0.00'
    for j in range(15, 19): A_.cell(i, j).fill = fill(INPUT)
    for j in range(1, 15): A_.cell(i, j).font = F()
    A_.cell(i, 9).alignment = Alignment(wrap_text=True, vertical='top')
dv = DataValidation(type='list', formula1='"Keep,Change,Remove"', allow_blank=True); A_.add_data_validation(dv); dv.add('O5:O%d' % (4 + len(order)))
A_.freeze_panes = 'D5'; A_.auto_filter.ref = 'A4:R%d' % (4 + len(order))

# ---------------- permanent or temporary ----------------
E_ = wb.create_sheet('Permanent or temporary')
title(E_, 'How each customer\'s entries apply', 'Rolling entries stay in force until the next entry for the account (0% ends one). Months-only entries apply to the months entered. Effect is the whole customer, GWh.')
header(E_, 4, ['KAM', 'Customer', 'How it applies', 'Reason', 'From', 'Last month entered', 'Adjustment', 'Effect FY2026', 'Effect FY2027', 'Effect FY2028', 'Permanent? (Yes / No / Until when)', 'Comment'], [20, 36, 20, 24, 10, 14, 34, 12, 12, 12, 24, 40])
r_ = 5
for kind, lst in (('Rolls forward', ROLL_CUST), ('Months entered only', OP_CUST)):
    for c in lst:
        E_.cell(r_, 1, c['kam']); E_.cell(r_, 2, c['name']); E_.cell(r_, 3, kind); E_.cell(r_, 4, ', '.join(sorted(c['reasons']))); E_.cell(r_, 5, t_label(c['first'])); E_.cell(r_, 6, t_label(c['last']))
        E_.cell(r_, 7, '; '.join(c['items'][:2]))
        for j in range(3): E_.cell(r_, 8 + j, c['eff'][j]).number_format = '+#,##0.00;-#,##0.00;0.00'
        E_.cell(r_, 11).fill = fill(INPUT); E_.cell(r_, 12).fill = fill(INPUT)
        for j in range(1, 11): E_.cell(r_, j).font = F()
        r_ += 1

# ---------------- new customers ----------------
N_ = wb.create_sheet('New customers')
title(N_, 'New customers in the forecast', 'Confirm each one. Volume is kWh a month from the start date.')
header(N_, 4, ['Customer', 'KAM', 'Rate class', 'Parish', 'Industry', 'Start', 'kWh a month', 'Entered by', 'Confirm / Change / Remove', 'Comment'], [34, 20, 10, 16, 26, 12, 14, 18, 22, 40])
for i, p in enumerate(PIPE, 5):
    for j, v in enumerate([p['name'], p['kam'] or 'Not assigned', p['rc'], p['parish'], p['ind'], p['start'], p['kwh'], p['by']], 1):
        N_.cell(i, j, v).font = F()
    N_.cell(i, 7).number_format = '#,##0'
    N_.cell(i, 9).fill = fill(INPUT); N_.cell(i, 10).fill = fill(INPUT)
dv2 = DataValidation(type='list', formula1='"Confirm,Change,Remove"', allow_blank=True); N_.add_data_validation(dv2); dv2.add('I5:I%d' % (4 + len(PIPE)))

# ---------------- account forecast ----------------
X = wb.create_sheet('Account forecast')
title(X, 'Account forecast, MWh by month', 'October 2026 to December 2028, with full-year totals. FY2026 includes January to September actual.')
months = [(2026, m) for m in range(10, 13)] + [(y, m) for y in (2027, 2028) for m in range(1, 13)]
header(X, 4, ['KAM', 'Account', 'Customer', 'Rate class', 'Industry', 'FY2026', 'FY2027', 'FY2028'] + ['%s %s' % (MON[m - 1], str(y)[2:]) for y, m in months], [20, 17, 34, 10, 26, 11, 11, 11] + [9] * len(months))
i = 5
for a in sorted(accts.values(), key=lambda a: (a['kam'] or 'zzz', -fy(a, '2027'))):
    X.cell(i, 1, a['kam'] or 'Not assigned'); X.cell(i, 2, a['ac']); X.cell(i, 3, a['name']); X.cell(i, 4, a['rc']); X.cell(i, 5, a['ind'])
    for j, y in enumerate(('2026', '2027', '2028')): X.cell(i, 6 + j, fy(a, y) * 1000).number_format = '#,##0.0'
    for j, (y, m) in enumerate(months): X.cell(i, 9 + j, a['live'][str(y)][m - 1] * 1000).number_format = '#,##0.0'
    i += 1
X.freeze_panes = 'D5'; X.auto_filter.ref = 'A4:%s%d' % (L(8 + len(months)), i - 1)

out = r'C:\Projects\Sales_Platform\JPS_KAM_Forecast_Round_Workbook_Oct2026_v2.xlsx'
wb.save(out); print('saved', out, 'accounts', i - 5, 'adjustments', len(order), 'miss rows', len(rows), 'customers', sum(1 for _ in CUST))
