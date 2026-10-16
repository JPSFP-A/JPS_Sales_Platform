# -*- coding: utf-8 -*-
import collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
import compute as c, engine as e

rows = sorted(c.rows, key=lambda r: -sum(r['ch'][m]['total'] for m in range(1, 13)))
NAVY, CYAN, LIGHT, GREY = '062552', '009FDA', 'D9E1F2', '66788A'
def F(**k):
    return Font(name='Arial', size=k.pop('size', 10), **k)
def hfill(col):
    return PatternFill('solid', fgColor=col)
MON = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
CLASSES = ['RT20', 'RT40', 'RT50', 'RT60-ST', 'RT70']
wb = Workbook()

# ---------------- Assumptions ----------------
A = wb.active; A.title = 'Assumptions'
A['A1'] = 'Forecast assumptions used (Oct-Dec 2026), read from the Sales Platform tables'; A['A1'].font = F(bold=True, size=13, color=NAVY)
A['A3'] = 'All-in rate by class, J$/kWh (jps_rate_assumptions: base + energy + demand + customer charge)'; A['A3'].font = F(bold=True)
for j, h in enumerate(['Rate class', 'Oct', 'Nov', 'Dec']):
    x = A.cell(4, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = hfill(NAVY)
for i, rc in enumerate(CLASSES):
    A.cell(5 + i, 1, rc).font = F()
    for j, m in enumerate((10, 11, 12)):
        x = A.cell(5 + i, 2 + j, round(c.RATE[(rc, m)], 4)); x.number_format = '#,##0.0000'; x.font = F(color='0000FF')
A['A12'] = 'Percentage drivers applied to the rolling 3-month base (jps_macro_assumptions)'; A['A12'].font = F(bold=True)
for j, h in enumerate(['Driver', 'Segment', 'Oct %', 'Nov %', 'Dec %']):
    x = A.cell(13, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = hfill(NAVY)
r_ = 14
for t, s in [('macro', 'residential'), ('macro', 'commercial'), ('weather', None)]:
    A.cell(r_, 1, t).font = F(); A.cell(r_, 2, s or 'all').font = F()
    for j, m in enumerate((10, 11, 12)):
        x = A.cell(r_, 3 + j, e.macro(t, s, 2026, m)); x.number_format = '0.00'; x.font = F(color='0000FF')
    r_ += 1
A.cell(r_, 1, 'RT20 uses the residential macro row; RT40/50/60-ST/70 use commercial. Industry driver applies to non-RT20 accounts with an industry tag.').font = F(color=GREY, italic=True)
r_ += 2
A.cell(r_, 1, 'Industry drivers (Oct-Dec 2026)').font = F(bold=True); r_ += 1
for j, h in enumerate(['Industry', '', 'Oct %', 'Nov %', 'Dec %']):
    x = A.cell(r_, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = hfill(NAVY)
r_ += 1
inds = sorted({k[1] for k in e.MAC if k[0] == 'industry' and k[2] == 2026 and k[3] >= 10})
for ind in inds:
    A.cell(r_, 1, ind).font = F()
    for j, m in enumerate((10, 11, 12)):
        x = A.cell(r_, 3 + j, e.macro('industry', ind, 2026, m)); x.number_format = '0.00'; x.font = F(color='0000FF')
    r_ += 1
r_ += 1
A.cell(r_, 1, 'Operational adjustments (jps_operational_adjustments_current), Oct-Dec 2026').font = F(bold=True); r_ += 1
for j, h in enumerate(['Account', 'Month', 'Adj %', 'Manual kWh', 'Reason']):
    x = A.cell(r_, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = hfill(NAVY)
r_ += 1
for o in sorted([o for o in e.D['opadj'] if o['year'] == 2026 and o['month'] >= 10], key=lambda o: (o['jps_ac'], o['month'])):
    for j, v in enumerate([o['jps_ac'], MON[o['month'] - 1], o['operational_pct'], o['manual_kwh'], o.get('reason_code')]):
        A.cell(r_, 1 + j, v).font = F()
    r_ += 1
A.column_dimensions['A'].width = 52; A.column_dimensions['B'].width = 14
for k in 'CDE':
    A.column_dimensions[k].width = 14

# ---------------- Customers ----------------
S = wb.create_sheet('Customers')
H = ['Account', 'Customer', 'Rate class', 'Parish', 'Industry', 'Managed?', 'KAM', 'Record type'] \
    + [m + ' kWh' for m in MON] + ['YTD kWh (Jan-Sep)', 'Q4 kWh (Oct-Dec fcst)', 'FY2026 kWh'] \
    + [m + ' kVA' for m in MON] + ['FY2026 kVA (sum of months)', 'Peak month kVA'] \
    + ['YTD revenue J$ (Jan-Sep actual)', 'Q4 revenue J$ (fcst)', 'FY2026 revenue J$', 'Notes']
S['A1'] = 'JPS commercial customers: FY2026 kWh, kVA and revenue (Jan-Sep actual, Oct-Dec Sales Platform Driver Forecast)'; S['A1'].font = F(bold=True, size=13, color=NAVY)
S['A2'] = 'Unmanaged accounts carry Jordache Wilson as KAM and are forecast at account level. Revenue excludes GCT, as in jps_actuals. RT10 residential is not account-level in the platform and is not in this file.'; S['A2'].font = F(color=GREY, italic=True)
HR = 4; R0 = HR + 1
col = {h: i + 1 for i, h in enumerate(H)}
for h, i in col.items():
    x = S.cell(HR, i, h); x.font = F(bold=True, color='FFFFFF')
    fc = (h.split()[0] in ('Oct', 'Nov', 'Dec') and ('kWh' in h or 'kVA' in h)) or h.startswith('Q4')
    x.fill = hfill(CYAN if fc else NAVY); x.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
S.row_dimensions[HR].height = 42
cK = col['Jan kWh']; cY, cQ, cF = col['YTD kWh (Jan-Sep)'], col['Q4 kWh (Oct-Dec fcst)'], col['FY2026 kWh']
cV = col['Jan kVA']; cVF, cVP = col['FY2026 kVA (sum of months)'], col['Peak month kVA']
cR, cRQ, cRF, cN = col['YTD revenue J$ (Jan-Sep actual)'], col['Q4 revenue J$ (fcst)'], col['FY2026 revenue J$'], col['Notes']
opnote = collections.defaultdict(list)
for o in e.D['opadj']:
    if o['year'] == 2026 and o['month'] >= 10 and (o['operational_pct'] or o['manual_kwh']):
        opnote[o['jps_ac']].append('%s %+g%%' % (MON[o['month'] - 1], o['operational_pct']))
for i, r in enumerate(rows):
    n = R0 + i
    S.cell(n, 1, r['ac']); S.cell(n, 2, r['name']); S.cell(n, 3, r['rc']); S.cell(n, 4, r['parish']); S.cell(n, 5, r['industry'])
    S.cell(n, 6, 'Yes' if r['managed'] else 'No'); S.cell(n, 7, r['kam']); S.cell(n, 8, r['kind'])
    for m in range(1, 13):
        S.cell(n, cK + m - 1, round(r['ch'][m]['total'], 2))
    S.cell(n, cY, '=SUM(%s%d:%s%d)' % (L(cK), n, L(cK + 8), n))
    S.cell(n, cQ, '=SUM(%s%d:%s%d)' % (L(cK + 9), n, L(cK + 11), n))
    S.cell(n, cF, '=%s%d+%s%d' % (L(cY), n, L(cQ), n))
    if r['has_kva']:
        for m in range(1, 13):
            S.cell(n, cV + m - 1, round(r['kva'][m], 1))
        S.cell(n, cVF, '=SUM(%s%d:%s%d)' % (L(cV), n, L(cV + 11), n))
        S.cell(n, cVP, '=MAX(%s%d:%s%d)' % (L(cV), n, L(cV + 11), n))
    S.cell(n, cR, round(r['rev_ytd'], 2))
    q = '+'.join('%s%d*INDEX(Assumptions!$%s$5:$%s$9,MATCH($C%d,Assumptions!$A$5:$A$9,0))' % (L(cK + 9 + j), n, L(2 + j), L(2 + j), n) for j in range(3))
    S.cell(n, cRQ, '=' + q)
    S.cell(n, cRF, '=%s%d+%s%d' % (L(cR), n, L(cRQ), n))
    notes = []
    if r['ac'] in opnote:
        notes.append('Operational adj: ' + ', '.join(opnote[r['ac']]))
    if r['kind'] == 'Bulk (no account)':
        notes.append('Bulk-billed lines with no individual account; forecast as a cohort')
    if r['kind'].startswith('New'):
        notes.append('Pipeline customer from jps_forecast_new_customers')
    if r['ac'] == '100185-607213':
        notes.append('Alcoa: Aug/Sep kVA and demand charge provisional, to be adjusted down')
    if notes:
        S.cell(n, cN, '; '.join(notes))
last = R0 + len(rows) - 1
for rr_ in S.iter_rows(min_row=R0, max_row=last):
    for x in rr_:
        x.font = F()
        if x.column >= cK and x.column != cN:
            x.number_format = '#,##0;(#,##0);-'
S.freeze_panes = S.cell(R0, 3); S.auto_filter.ref = 'A%d:%s%d' % (HR, L(cN), last)
for k, w in {'A': 18, 'B': 38, 'C': 9, 'D': 14, 'E': 30, 'F': 10, 'G': 20, 'H': 18}.items():
    S.column_dimensions[k].width = w
for j in range(cK, cRF + 1):
    S.column_dimensions[L(j)].width = 13
S.column_dimensions[L(cN)].width = 60

# ---------------- Summary ----------------
M = wb.create_sheet('Summary', 0)
M['A1'] = 'FY2026 commercial customer forecast: summary by rate class and management status'; M['A1'].font = F(bold=True, size=13, color=NAVY)
M['A2'] = 'All figures are formulas over the Customers sheet. kWh in GWh, revenue in J$ million, kVA as the sum of monthly billed kVA.'; M['A2'].font = F(color=GREY, italic=True)
hd = ['Rate class', 'Managed?', 'Accounts', 'YTD GWh (Jan-Sep)', 'Q4 GWh (fcst)', 'FY2026 GWh', 'FY2026 kVA (sum of months)', 'YTD revenue J$M', 'Q4 revenue J$M (fcst)', 'FY2026 revenue J$M']
for j, h in enumerate(hd):
    x = M.cell(4, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = hfill(NAVY); x.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
M.row_dimensions[4].height = 42
def rng(cc):
    return 'Customers!$%s$%d:$%s$%d' % (L(cc), R0, L(cc), last)
n = 5; first = n
for rc in CLASSES:
    for mg in ('Yes', 'No'):
        M.cell(n, 1, rc); M.cell(n, 2, mg)
        crit = '%s,$A%d,%s,$B%d' % (rng(3), n, rng(6), n)
        M.cell(n, 3, '=COUNTIFS(%s)' % crit)
        for j, (cc, div) in enumerate([(cY, 1000000), (cQ, 1000000), (cF, 1000000), (cVF, 1), (cR, 1000000), (cRQ, 1000000), (cRF, 1000000)]):
            M.cell(n, 4 + j, '=SUMIFS(%s,%s)' % (rng(cc), crit) + ('/%d' % div if div != 1 else ''))
        n += 1
last_s = n - 1
M.cell(n, 1, 'Total')
for j in range(3, 11):
    M.cell(n, j, '=SUM(%s%d:%s%d)' % (L(j), first, L(j), last_s))
for rr_ in M.iter_rows(min_row=5, max_row=n):
    for x in rr_:
        x.font = F(bold=(rr_[0].row == n))
        if x.column >= 3:
            x.number_format = '#,##0.0' if x.column not in (3, 7) else '#,##0'
        if rr_[0].row == n:
            x.fill = hfill(LIGHT)
n += 2
for t in ['Managed = account (or its customer prefix) has a KAM in jps_kam. Unmanaged accounts are shown with Jordache Wilson as KAM.',
          'kVA exists only for demand-metered accounts (RT40, RT50, RT60-ST, RT70). RT20 is energy-only, so its kVA is blank.']:
    M.cell(n, 1, t).font = F(color=GREY, italic=True); n += 1
M.column_dimensions['A'].width = 12; M.column_dimensions['B'].width = 11
for j in range(3, 11):
    M.column_dimensions[L(j)].width = 16

# ---------------- Method & Recon ----------------
R = wb.create_sheet('Method & Recon')
lines = [('Method', True),
 ('Source: Sales Platform Driver Forecast engine (_dfcComputeChain), ported line for line. Closed through September (Jan-Sep actual from jps_actuals).', False),
 ('Oct-Dec kWh = rolling average of the prior three months (forecast months roll in) x (1 + macro + weather + industry + operational %) + manual kWh. The percentage sum is floored at -100%.', False),
 ('Industry driver applies to non-RT20 accounts with an industry tag. RT20 uses the residential macro and no industry driver, as in the platform.', False),
 ('Pipeline customers (jps_forecast_new_customers) with a 2026 start are included. 2027 starters are not.', False),
 ('kVA: the platform has no account-level kVA forecast. Jan-Sep is billed kVA from the monthly scan. Oct-Dec = prior three-month kVA average x the same driver factor as that account kWh (forecast kWh / base kWh).', False),
 ('Revenue: Jan-Sep actual revenue_jmd (net of GCT). Oct-Dec = forecast kWh x the platform all-in class rate (Assumptions sheet). It is a class rate, not an account tariff, so account-level revenue is indicative.', False),
 ('Unmanaged accounts: the platform forecasts these only as cohorts (class | parish | bucket). Here they run as individual accounts with Jordache Wilson as KAM.', False),
 ('', False), ('Reconciliation (Q4 forecast, GWh)', True)]
for i, (t, b) in enumerate(lines, 1):
    R.cell(i, 1, t).font = F(bold=b, color=NAVY if b else '000000')
rr = len(lines) + 1
for j, h in enumerate(['Rate class', 'This file', 'Platform as built (KAM accounts + cohorts)', 'Difference MWh']):
    x = R.cell(rr, 1 + j, h); x.font = F(bold=True, color='FFFFFF'); x.fill = hfill(NAVY); x.alignment = Alignment(wrap_text=True)
rr += 1; r1 = rr
for rc in CLASSES:
    R.cell(rr, 1, rc).font = F()
    R.cell(rr, 2, '=SUMIFS(%s,%s,A%d)/1000000' % (rng(cQ), rng(3), rr)).font = F()
    x = R.cell(rr, 3, round(sum(c.asis[rc].values()) / 1e6, 4)); x.font = F(color='0000FF')
    R.cell(rr, 4, '=(B%d-C%d)*1000' % (rr, rr)).font = F()
    rr += 1
R.cell(rr, 1, 'Total').font = F(bold=True)
for j in (2, 3, 4):
    R.cell(rr, j, '=SUM(%s%d:%s%d)' % (L(j), r1, L(j), rr - 1)).font = F(bold=True)
for row in R.iter_rows(min_row=r1, max_row=rr, min_col=2, max_col=4):
    for x in row:
        x.number_format = '#,##0.000' if x.column < 4 else '#,##0.0'
rr += 2
for t in ['The difference is the industry driver applying to unmanaged RT40/RT50 accounts, which the platform runs at cohort level with no industry tag. It is the effect of making you the KAM.',
          'YTD tie to jps_actuals (Jan-Sep 2026, RT20-RT70): 1,546.436 GWh in the database. Differences of a few MWh by class come from accounts reclassified mid-year, which the platform books to the latest class.',
          'Accounts with no 2026 billing (%s billed in 2025 only) are left out; they forecast at zero.' % format(len(c.skipped_2025only), ',')]:
    R.cell(rr, 1, t).font = F(color=GREY, italic=True); rr += 1
R.column_dimensions['A'].width = 150
for k in 'BCD':
    R.column_dimensions[k].width = 22
wb.move_sheet('Assumptions', offset=3)
out = r'C:\Projects\Sales_Platform\JPS_Customer_Forecast_FY2026_by_Account.xlsx'
wb.save(out); print('saved', out, len(rows))
