# -*- coding: utf-8 -*-
# Redraft of the Client Relations and Commercial Update (Sep 2026): corrections, month and YTD on the sales slide,
# a revenue and retention slide, and an outlook and recovery slide. Edits the original file in place (saved under a new name).
import copy, json, collections, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK

SRC = 'cr_orig.pptx'
OUT = r'C:\Projects\Sales_Platform\Client Relations and Commercial Update (Sep2026)_v4.pptx'
NAVY, BLUE, CYAN, GREY, LIGHT = RGBColor(0x06, 0x25, 0x52), RGBColor(0x0E, 0x4E, 0x95), RGBColor(0x00, 0x9F, 0xDA), RGBColor(0x66, 0x78, 0x8A), RGBColor(0xD9, 0xE1, 0xF2)
RED, GREEN, WHITE, BLACK = RGBColor(0xC0, 0x00, 0x00), RGBColor(0x1F, 0x8F, 0x5A), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x1F, 0x1F, 0x1F)
F_GOOD, F_MID, F_BAD, F_ALT, F_TILE = RGBColor(0xDD, 0xEF, 0xDF), RGBColor(0xFF, 0xF0, 0xC8), RGBColor(0xF6, 0xD3, 0xD6), RGBColor(0xF3, 0xF6, 0xFB), RGBColor(0xEE, 0xF2, 0xF9)
FONT = 'Arial'
MINUS = '\u2212'

# ---------------- verified numbers ----------------
CLASS = ['RT20', 'RT40', 'RT50', 'RT60-ST', 'RT70']
NAMES = {'RT20': 'RT20 Small commercial', 'RT40': 'RT40 Commercial', 'RT50': 'RT50 Large commercial', 'RT60-ST': 'RT60 Street lighting', 'RT70': 'RT70 Industrial'}
# GWh: sep actual, sep budget, sep 2025, ytd actual, ytd budget, ytd 2025
D = {'RT20': (60.618, 57.649, 58.278, 475.762, 461.306, 490.521), 'RT40': (67.313, 73.334, 68.941, 569.767, 584.913, 605.324),
     'RT50': (30.253, 35.268, 33.527, 259.549, 274.111, 287.170), 'RT60-ST': (3.413, 3.312, 3.409, 27.490, 27.057, 30.382), 'RT70': (18.643, 22.866, 27.281, 213.868, 192.712, 226.863)}
ABS = {'RT40': (22.938, 106.260), 'RT50': (10.841, 60.604), 'RT60-ST': (0.138, 0.726), 'RT70': (13.490, 45.296)}   # sum of account misses: sep, ytd
def acc(rc, per):
    if rc not in ABS: return None
    a = D[rc][0] if per == 0 else D[rc][3]
    return 1 - ABS[rc][per] / a
def pct(x, d=1):
    s = f'{abs(x) * 100:.{d}f}%'
    return ('+' if x > 0.00004 else (MINUS if x < -0.00004 else '')) + s
def sg(x, d=1): return ('+' if x > 0 else (MINUS if x < 0 else '')) + f'{abs(x):,.{d}f}'
tot = [sum(D[c][i] for c in CLASS) for i in range(6)]
assert abs(tot[0] - 180.240) < 0.01 and abs(tot[1] - 192.429) < 0.01 and abs(tot[3] - 1546.436) < 0.01
abs_tot = [sum(v[i] for v in ABS.values()) for i in (0, 1)]
act_tot = [sum(D[c][i] for c in ABS) for i in (0, 3)]
acc_tot = [1 - abs_tot[0] / act_tot[0], 1 - abs_tot[1] / act_tot[1]]
IND = {'Mining (incl. Alcoa)': ((3.038, 10.258), (80.756, 83.915)), 'Hotels and motels': ((17.231, 20.944), (145.540, 156.689)), 'Water, irrigation': ((18.623, 21.042), (158.029, 166.537)),
       'Poultry': ((1.993, 4.105), (30.413, 33.056)), 'Telecoms': ((6.972, 8.189), (60.966, 65.744)), 'Cement': ((9.150, 8.728), (82.527, 75.219)),
       'Hospitals': ((4.433, 3.688), (33.461, 29.827)), 'Food manufacturing': ((4.407, 3.592), (34.282, 29.939)), 'Distilleries': ((3.462, 2.588), (28.163, 22.897))}

# key-account revenue (RT40 to RT70, Jan to Sep)
rows = json.load(open('keyacct_components.json'))
fl = lambda r, k: float(r.get(k) or 0)
A = {2025: collections.defaultdict(lambda: collections.defaultdict(float)), 2026: collections.defaultdict(lambda: collections.defaultdict(float))}
nm = {}
for r in rows:
    if not r['jps_ac']: continue
    a = A[r['year']][r['jps_ac']]
    a['kwh'] += fl(r, 'kwh'); a['nf'] += fl(r, 'energy_jmd') + fl(r, 'demand_jmd') + fl(r, 'customer_charge_jmd'); a['dem'] += fl(r, 'demand_jmd'); nm[r['jps_ac']] = r['name']
T_ = lambda y, k: sum(v[k] for v in A[y].values())
act25 = {ac for ac, v in A[2025].items() if v['kwh'] > 0}; act26 = {ac for ac, v in A[2026].items() if v['kwh'] > 0}
logo = len(act25 & act26) / len(act25)
def retention(k):
    base = sum(A[2025][ac][k] for ac in act25)
    return sum(A[2026][ac][k] for ac in act25) / base, sum(min(A[2026][ac][k], A[2025][ac][k]) for ac in act25) / base
vol_net, vol_gross = retention('kwh'); nf_net, nf_gross = retention('nf')
# customer-name level view of the key accounts
byname = collections.defaultdict(lambda: {2025: collections.defaultdict(float), 2026: collections.defaultdict(float)})
for y in (2025, 2026):
    for ac, v in A[y].items():
        for k in ('kwh', 'nf'): byname[nm[ac].strip()][y][k] += v[k]
top8 = sorted(byname.items(), key=lambda kv: -kv[1][2026]['nf'])
top8 = [kv for kv in top8 if kv[1][2026]['kwh'] > 0][:8]
moves = sorted(((v[2026]['nf'] - v[2025]['nf'], n) for n, v in byname.items()))
dec6, inc5 = moves[:6], list(reversed(moves[-5:]))

# monthly year-on-year (billed GWh), commercial total, hotels, water
d = json.load(open('pull.json')); nn = json.load(open('pull_nc.json'))
ind = {str(k['jps_ac']).strip(): k['industry'] or '' for k in d['ind']}
kl = lambda ac: ind.get(ac) or ind.get(str(ac or '').split('-')[0]) or ''
com = collections.defaultdict(float); hot = collections.defaultdict(float); wat = collections.defaultdict(float)
for r in d['act']:
    k = float(r['kwh'] or 0); y, m = r['year'], r['month']
    com[(y, m)] += k
    i = kl(r['jps_ac']) if r['jps_ac'] else ''
    if 'Hotel' in i or 'Motel' in i: hot[(y, m)] += k
    if 'Water Supply' in i: wat[(y, m)] += k
yoy = lambda dd: [dd[(2026, m)] / dd[(2025, m)] - 1 for m in range(1, 10)]
Y_COM, Y_HOT, Y_WAT = yoy(com), yoy(hot), yoy(wat)
assert abs(sum(com[(2026, m)] for m in range(1, 10)) / 1e6 - 1546.436) < 0.05

prs = Presentation(SRC)

# ---------------- helpers ----------------
def set_text(shape, new):
    tf = shape.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    runs[0].text = new
    for r in runs[1:]: r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
def by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid: return sh
    raise KeyError(sid)
def tb(slide, x, y, w, h, text, size=11, bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h))); tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0); tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        for t, b, c in (ln if isinstance(ln, list) else [(ln, bold, color)]):
            r = p.add_run(); r.text = t; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = c
    return s
def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s
def cell_set(c, text, size=9, bold=False, color=BLACK, align=PP_ALIGN.CENTER, fill=None):
    c.text = ''; tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    c.margin_left = c.margin_right = Emu(36000); c.margin_top = c.margin_bottom = Emu(12000); c.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is None: c.fill.background()
    else: c.fill.solid(); c.fill.fore_color.rgb = fill
def table(slide, x, y, w, colw, rows_, row_h, size=9, head_rows=1, left=(0,), colors=None, total=False):
    nr, nc = len(rows_), len(rows_[0])
    gf = slide.shapes.add_table(nr, nc, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(row_h * nr))); t = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for a in ('bandRow', 'firstRow'): tblPr.set(a, '0')
    sid = tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
    if sid is not None: sid.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'
    s_ = sum(colw)
    for j, cw in enumerate(colw): t.columns[j].width = Emu(int(w * cw / s_))
    for i in range(nr): t.rows[i].height = Emu(int(row_h))
    for i, row in enumerate(rows_):
        hdr = i < head_rows; last = total and i == nr - 1
        for j, v in enumerate(row):
            if hdr: cell_set(t.cell(i, j), v, size, True, WHITE, PP_ALIGN.LEFT if j in left else PP_ALIGN.CENTER, NAVY)
            else:
                fill = F_ALT if last else None
                col = colors[i - head_rows][j] if colors and colors[i - head_rows][j] is not None else (NAVY if j == 0 else BLACK)
                cell_set(t.cell(i, j), v, size, last or j == 0, col, PP_ALIGN.LEFT if j in left else PP_ALIGN.CENTER, fill)
    return t
def clone(src_shape, dst_slide):
    el = copy.deepcopy(src_shape._element); dst_slide.shapes._spTree.append(el); return dst_slide.shapes[-1]
def new_from(slide5, title, sub, init_head, init_text):
    layout = slide5.slide_layout
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    t = clone(by_id(slide5, 2), s); set_text(t, title)
    sb = clone(by_id(slide5, 3), s); set_text(sb, sub)
    ki = clone(by_id(slide5, 21), s); ch = clone(by_id(slide5, 22), s)
    # key initiative text: first run heading, second run body (match the original two-run layout)
    p = ki.text_frame.paragraphs[0]; rr = p.runs
    return s, ki, rr
col_ = lambda v: RED if v < 0 else GREEN
S1, S5 = prs.slides[0], prs.slides[4]

# ================= slide 1 =================
set_text(by_id(S1, 3), 'Service on target | Commercial sales 6.3% below budget, J$4.16B at risk: building a commercial model')
set_text(by_id(S1, 51), 'Sales')
set_text(by_id(S1, 54), 'Sept commercial sales vs budget: 180.2 vs 192.4 GWh')
set_text(by_id(S1, 60), f'{acc_tot[0] * 100:.0f}% Sept, {acc_tot[1] * 100:.0f}% YTD')
set_text(by_id(S1, 57), f'{pct(tot[3] / tot[4] - 1)} bud, {pct(tot[3] / tot[5] - 1)} 2025')
set_text(by_id(S1, 56), 'YTD')
set_text(by_id(S1, 39), 'Lost')
set_text(by_id(S1, 40), 'Alcoa (Aug), PAC (Oct)')
set_text(by_id(S1, 71), "Alcoa's move to self-generation explains 6.5 GWh of the 12.2 GWh September gap")

for sid in range(12, 25): by_id(S1, sid)._element.getparent().remove(by_id(S1, sid)._element)
ch1 = by_id(S1, 11); ch1.height = Emu(1500000)
by_id(S1, 25).top = Emu(4500000)

# ================= slide 5 =================
set_text(by_id(S5, 3), "Commercial sales 12.2 GWh (6.3%) below budget; Alcoa's move to self-generation explains 6.5 GWh")
set_text(by_id(S5, 4), 'Commercial sales by rate class, September and January to September (GWh)')
old = by_id(S5, 5); x, y, w = old.left, old.top, old.width
old._element.getparent().remove(old._element)
head1 = ['', 'September', '', '', '', 'January to September', '', '', '', '']
head2 = ['Rate class', 'Actual', 'Budget', 'vs bud', 'vs 2025', 'Actual', 'Budget', 'vs bud', 'vs 2025', 'Acc']
body, colors = [], []
for rc in CLASS:
    a, b, p, ya, yb, yp = D[rc]
    ac0, ac1 = acc(rc, 0), acc(rc, 1)
    body.append([NAMES[rc], f'{a:.1f}', f'{b:.1f}', pct(a / b - 1), pct(a / p - 1), f'{ya:.1f}', f'{yb:.1f}', pct(ya / yb - 1), pct(ya / yp - 1), 'n/a' if ac1 is None else f'{ac1 * 100:.0f}%'])
    colors.append([None, None, None, col_(a / b - 1), col_(a / p - 1), None, None, col_(ya / yb - 1), col_(ya / yp - 1), None])
body.append(['Commercial total', f'{tot[0]:.1f}', f'{tot[1]:.1f}', pct(tot[0] / tot[1] - 1), pct(tot[0] / tot[2] - 1), f'{tot[3]:,.1f}', f'{tot[4]:,.1f}', pct(tot[3] / tot[4] - 1), pct(tot[3] / tot[5] - 1), f'{acc_tot[1] * 100:.0f}%*'])
colors.append([None, None, None, col_(tot[0] / tot[1] - 1), col_(tot[0] / tot[2] - 1), None, None, col_(tot[3] / tot[4] - 1), col_(tot[3] / tot[5] - 1), None])
t = table(S5, x, y, w, [1.9, 0.7, 0.7, 0.75, 0.75, 0.85, 0.85, 0.75, 0.75, 0.6], [head1, head2] + body, 236000, size=8, head_rows=2, left=(0,), colors=colors, total=True)
t.cell(0, 1).merge(t.cell(0, 4)); t.cell(0, 5).merge(t.cell(0, 9))
cell_set(t.cell(0, 1), 'September', 8, True, WHITE, PP_ALIGN.CENTER, NAVY); cell_set(t.cell(0, 5), 'January to September', 8, True, WHITE, PP_ALIGN.CENTER, NAVY)
for i, rc in enumerate(CLASS, 2):
    for j, per in ((9, 1),):
        v = acc(rc, per)
        if v is not None:
            c = t.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = F_GOOD if v >= 0.95 else (F_MID if v >= 0.85 else F_BAD)
set_text(by_id(S5, 6), 'Acc: YTD account budget vs billing, RT40-RT70. Jan-Feb budget = actuals, so YTD vs budget flatters.')
# industry chart: September and YTD variance
oldc = by_id(S5, 8); cx, cy, cw, ch = oldc.left, oldc.top, oldc.width, oldc.height
oldc._element.getparent().remove(oldc._element)
cats = list(reversed(list(IND.keys())))
cd = CategoryChartData(); cd.categories = cats
cd.add_series('YTD', [IND[k][1][0] - IND[k][1][1] for k in cats]); cd.add_series('September', [IND[k][0][0] - IND[k][0][1] for k in cats])
gf = S5.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, cx, cy, cw, ch, cd); chart = gf.chart
chart.has_title = False; chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout = False; chart.legend.font.size = Pt(8); chart.legend.font.name = FONT
plot = chart.plots[0]; plot.gap_width = 40; plot.overlap = -5
for ser, c_ in zip(plot.series, (CYAN, NAVY)): ser.format.fill.solid(); ser.format.fill.fore_color.rgb = c_; ser.invert_if_negative = False
plot.has_data_labels = True; dl = plot.data_labels; dl.number_format = '+0.0;-0.0'; dl.number_format_is_linked = False; dl.font.size = Pt(7); dl.font.name = FONT; dl.position = XL_LABEL_POSITION.OUTSIDE_END
ca = chart.category_axis; ca.tick_labels.font.size = Pt(8); ca.tick_labels.font.name = FONT; ca.has_major_gridlines = False; ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
va = chart.value_axis; va.has_major_gridlines = False; va.tick_label_position = XL_TICK_LABEL_POSITION.NONE; va.major_tick_mark = XL_TICK_MARK.NONE
set_text(by_id(S5, 7), 'Variance vs budget by industry (GWh)')
# spotlight and commentary
set_text(by_id(S5, 15), 'up to J$175M')
sp = by_id(S5, 17)
p0 = sp.text_frame.paragraphs[0]
p0.runs[0].text = 'New turbine (TG4) commissioned in August: the refinery now generates all its own power; grid draw fell from 9.0 GWh in July to 0.6 in Sept; load not expected back'
for r in p0.runs[1:]: r._r.getparent().remove(r._r)
wm = by_id(S5, 20); wp = wm.text_frame.paragraphs
wp[2].runs[0].text = f'YTD: commercial sales {pct(tot[3] / tot[4] - 1)} vs budget but {pct(tot[3] / tot[5] - 1)} vs 2025; Sept shows the gap more sharply as Alcoa and hotels weigh on the month'
for r in wp[2].runs[1:]: r._r.getparent().remove(r._r)
wp[1]._p.getparent().remove(wp[1]._p)


# ================= growth slide: pipeline and new load only =================
S6 = prs.slides[5]
for sid in range(4, 22): by_id(S6, sid)._element.getparent().remove(by_id(S6, sid)._element)
set_text(by_id(S6, 2), 'COMMERCIAL  |  GROWTH & PIPELINE')
set_text(by_id(S6, 3), 'The 2027 pipeline is J$1.17B across 13 projects, led by NWC; 25.67 GWh of new load connects in 2026')
pn = by_id(S6, 22); pn.left = Emu(612648); pn.width = Emu(10972800); pn.height = Emu(4300000)
for sid in (23, 24, 25, 26):
    sh = by_id(S6, sid); sh.left = Emu(sh.left - 6199632)
by_id(S6, 27)._element.getparent().remove(by_id(S6, 27)._element)
for sid in (28, 29, 30, 31):
    sh = by_id(S6, sid); sh.left = Emu(sh.left - 795160); sh.top = Emu(sh.top - (1700784 if sid in (28, 29) else 1646512 if sid == 30 else 1150000))
for r in by_id(S6, 26).text_frame.paragraphs:
    for rn in r.runs: rn.font.size = Pt(14)
    r.space_after = Pt(14)
by_id(S6, 26).top = Emu(2500000)
tbl = by_id(S6, 30).table
for rw in tbl.rows:
    rw.height = Emu(450000)
    for c in rw.cells:
        for p in c.text_frame.paragraphs:
            for rn in p.runs: rn.font.size = Pt(11)
by_id(S6, 31).top = Emu(4850000)
for p in by_id(S6, 31).text_frame.paragraphs:
    for rn in p.runs: rn.font.size = Pt(11)
kg = by_id(S6, 32).text_frame.paragraphs[-1].runs
kg[0].text = 'Pipeline delivery: '
if len(kg) > 1: kg[1].text = "connect the six 2026 projects on time and secure NWC's 2027 load against its solar programme"
for r in kg[2:]: r._r.getparent().remove(r._r)

# ================= new slide: revenue and retention =================
SA, ki, rr = new_from(S5, 'COMMERCIAL  |  REVENUE & RETENTION', f'Key accounts keep {nf_net * 100:.0f}% of their revenue, but six accounts drive J${abs(sum(x for x, n in dec6)) / 1e6:,.0f}M of the J${abs(T_(2026, "nf") - T_(2025, "nf")) / 1e9:.1f}bn decline', '', '')
ki.text_frame.paragraphs[0].runs[0].text = 'Key Initiative'
body_runs = ki.text_frame.paragraphs[-1].runs
body_runs[0].text = 'Revenue-at-risk register: '
if len(body_runs) > 1: body_runs[1].text = 'price every Tier 1 account on non-fuel revenue and demand charge, not volume'
for r in body_runs[2:]: r._r.getparent().remove(r._r)
nf26, nf25 = T_(2026, 'nf'), T_(2025, 'nf'); k26, k25 = T_(2026, 'kwh'), T_(2025, 'kwh'); d26, d25 = T_(2026, 'dem'), T_(2025, 'dem')
def fixn(n):
    n = n.strip()
    if n.isupper(): n = n.title()
    n = n.replace('Restaurants Of Jam Ltd', 'Restaurants of Jamaica')
    n = re.sub(r'\s+(Limited|Limite|Limi|Ltd\.?|L)$', '', n)
    for a, b in [('Commissio', 'Commission'), ('Jamaican Reso', 'Jamaican Resorts'), ('Power Comp', 'Power Company'), ('Alcoa Mins Of Ja', 'Alcoa Minerals of Jamaica'), ('Caribbean Cement Co', 'Caribbean Cement'), (' Ja', ' Jamaica')]:
        if n.endswith(a): n = n[:-len(a)] + b
    return n
tiles = [(f'J${nf26 / 1e9:.1f}bn', 'non-fuel revenue, January to September', pct(nf26 / nf25 - 1)), (f'{k26 / 1e6:,.0f} GWh', 'volume, January to September', pct(k26 / k25 - 1))]
for i, (big, lab, chg) in enumerate(tiles):
    x0 = 612648 + i * 3250000
    rect(SA, x0, 1572768, 3150000, 880000, F_TILE)
    tb(SA, x0 + 120000, 1620000, 2900000, 330000, big, 20, True, NAVY)
    tb(SA, x0 + 120000, 1960000, 2900000, 250000, lab, 8.5, False, BLACK)
    tb(SA, x0 + 120000, 2200000, 2900000, 200000, chg + ' vs 2025', 9, True, col_(float(chg.replace('%', '').replace(MINUS, '-').replace('+', ''))))
tb(SA, 612648, 2560000, 6400000, 220000, 'Top key accounts by non-fuel revenue, January to September (J$M)', 10, True, NAVY)
rows_ = [['Account', '2026', '2025', 'Change', 'GWh 2026', 'GWh vs 2025']]
cols_ = []
for n, v in top8:
    a26, a25 = v[2026]['nf'] / 1e6, v[2025]['nf'] / 1e6
    rows_.append([fixn(n), f'{a26:,.0f}', f'{a25:,.0f}', pct(a26 / a25 - 1, 0) if a25 else 'new', f'{v[2026]["kwh"] / 1e6:,.1f}', sg(v[2026]['kwh'] / 1e6 - v[2025]['kwh'] / 1e6)])
    cols_.append([None, None, None, col_(a26 - a25), None, col_(v[2026]['kwh'] - v[2025]['kwh'])])
table(SA, 612648, 2800000, 6400000, [2.6, 0.8, 0.8, 0.8, 0.9, 1.0], rows_, 262000, size=8.5, left=(0,), colors=cols_)
# retention panel (top right)
rect(SA, 7315200, 1572768, 4270000, 1750000, None, LIGHT)
tb(SA, 7440000, 1632768, 4000000, 200000, 'RETENTION, JAN TO SEP 2026 VS 2025', 9, True, GREY)
ret = [('Accounts billing in both periods', f'{logo * 100:.1f}%', f'{len(act25 & act26):,} of {len(act25):,}'), ('Volume, net of growth', f'{vol_net * 100:.1f}%', f'{vol_gross * 100:.1f}% counting losses only'),
       ('Non-fuel revenue, net of growth', f'{nf_net * 100:.1f}%', f'{nf_gross * 100:.1f}% counting losses only')]
for i, (lab, big, sub) in enumerate(ret):
    yy = 1930000 + i * 450000
    tb(SA, 7440000, yy, 2300000, 400000, [[(lab, True, NAVY)], [(sub, False, GREY)]], 8.5)
    tb(SA, 9900000, yy, 1550000, 400000, big, 20, True, NAVY, PP_ALIGN.RIGHT)
tb(SA, 7315200, 3400000, 4270000, 500000, 'Counting accounts alone makes retention look strong. Volume is the better test: a single refinery moves it more than 100 small accounts.', 9, False, BLACK)
# movers: one list
tb(SA, 7315200, 4000000, 4270000, 200000, 'Largest changes in non-fuel revenue (J$M)', 9, True, NAVY)
mv = [(fixn(n), x / 1e6) for x, n in dec6] + [(fixn(n), x / 1e6) for x, n in inc5[:3]]
rows_m = [[nm_, (MINUS if v < 0 else '+') + f'{abs(v):,.0f}'] for nm_, v in mv]
table(SA, 7315200, 4220000, 4270000, [3.2, 1.0], [['Account', 'Change']] + rows_m, 172000, size=8, left=(0,), colors=[[None, col_(v)] for nm_, v in mv])
tb(SA, 612648, 5320000, 6400000, 700000, 'RT40 to RT70 accounts. Non-fuel revenue is energy, demand and customer charges; it excludes fuel, IPP and GCT. Net retention includes growth at retained accounts; the counting-losses-only figure caps each account at its 2025 level. Alcoa August and September demand charges are provisional.', 7.5, False, GREY)
SA.notes_slide.notes_text_frame.text = 'RT40 to RT70, January to September 2026 against 2025. Non-fuel revenue ties to billing components (energy, demand, customer charge). Retention is measured on accounts billed in both periods.'

# ================= new slide: outlook and recovery =================
com_cls = ['RT20', 'RT40', 'RT50', 'RT60-ST', 'RT70']
FY = {'RT20': (627.296, 654.672, 691.6, 708.4), 'RT40': (776.738, 771.913, 838.6, 835.3), 'RT50': (362.830, 345.651, 400.8, 433.4), 'RT60-ST': (38.647, 37.469, 39.8, 38.8), 'RT70': (288.124, 267.19, 193.7, 189.5)}
ctot = [sum(FY[c][i] for c in com_cls) for i in range(4)]
SB, ki2, rr2 = new_from(S5, 'COMMERCIAL  |  OUTLOOK & RECOVERY', f'Commercial grows {(ctot[2] / ctot[1] - 1) * 100:.1f}% in FY2027 as hotels recover; industrial falls {abs(FY["RT70"][2] / FY["RT70"][1] - 1) * 100:.0f}% to self-generation', '', '')
ki2.text_frame.paragraphs[0].runs[0].text = 'Key Initiative'
br = ki2.text_frame.paragraphs[-1].runs
br[0].text = 'Recovery tracking: '
if len(br) > 1: br[1].text = 'review hotel and water-supply volumes monthly against the reopening schedule'
for r in br[2:]: r._r.getparent().remove(r._r)
tb(SB, 612648, 1572768, 6000000, 220000, 'Commercial outlook by rate class (GWh)', 10, True, NAVY)
rows2 = [['Rate class', 'FY2025', 'FY2026', 'FY2027', 'YoY', 'FY2028', 'YoY']]; cols2 = []
for c in com_cls:
    v = FY[c]; rows2.append([NAMES[c], f'{v[0]:,.1f}', f'{v[1]:,.1f}', f'{v[2]:,.1f}', pct(v[2] / v[1] - 1), f'{v[3]:,.1f}', pct(v[3] / v[2] - 1)]); cols2.append([None, None, None, None, col_(v[2] - v[1]), None, col_(v[3] - v[2])])
rows2.append(['Commercial total', f'{ctot[0]:,.1f}', f'{ctot[1]:,.1f}', f'{ctot[2]:,.1f}', pct(ctot[2] / ctot[1] - 1), f'{ctot[3]:,.1f}', pct(ctot[3] / ctot[2] - 1)]); cols2.append([None, None, None, None, col_(ctot[2] - ctot[1]), None, col_(ctot[3] - ctot[2])])
table(SB, 612648, 1800000, 6200000, [2.4, 0.9, 0.9, 0.9, 0.8, 0.9, 0.8], rows2, 245000, size=8.5, left=(0,), colors=cols2, total=True)
rect(SB, 612648, 3650000, 6200000, 2200000, F_TILE)
tb(SB, 760000, 3700000, 5900000, 200000, 'WHAT MOVES THE OUTLOOK', 9, True, GREY)
tb(SB, 760000, 3910000, 5900000, 700000, [
    [('Recovery adds about 140 GWh in FY2027, ', True, NAVY), ('led by hotels returning in steps and water supply, which is nearly back to last year.', False, BLACK)],
    [('Defections take about 71 GWh. ', True, NAVY), ('Alcoa, MBJ Airports and PAC Kingston Airport are decisions already taken.', False, BLACK)]], 9.5)
tb(SB, 760000, 4650000, 5900000, 200000, 'HOTEL REOPENINGS: 5,018 ROOMS, J$4.52B A YEAR, 89.8% BACK BY MID-2027', 8.5, True, GREY)
table(SB, 760000, 4870000, 5900000, [1.6, 0.9, 1.0, 1.2], [['When', 'Hotels', 'Rooms', 'J$B a year'], ['December 2026', '6', '1,982', '1.78'], ['First half of 2027', '8', '2,524', '2.27'], ['2028', '1', '512', '0.46']], 235000, size=8.5, left=(0,))
tb(SB, 7100000, 1572768, 4500000, 220000, 'Year on year by month, 2026 vs 2025 (%)', 10, True, NAVY)
cdm = CategoryChartData(); cdm.categories = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep']
cdm.add_series('Commercial total', [v * 100 for v in Y_COM]); cdm.add_series('Hotels', [v * 100 for v in Y_HOT]); cdm.add_series('Water supply', [v * 100 for v in Y_WAT])
gf = SB.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Emu(7050000), Emu(1800000), Emu(4550000), Emu(3000000), cdm); chart = gf.chart
chart.has_title = False; chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout = False; chart.legend.font.size = Pt(9); chart.legend.font.name = FONT
for ser, c_ in zip(chart.plots[0].series, (NAVY, RED, CYAN)):
    ser.format.line.color.rgb = c_; ser.format.line.width = Pt(2.25); ser.smooth = False
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = c_; ser.marker.format.line.color.rgb = c_
ca = chart.category_axis; ca.tick_labels.font.size = Pt(9); ca.tick_labels.font.name = FONT; ca.has_major_gridlines = False; ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
va = chart.value_axis; va.tick_labels.font.size = Pt(9); va.tick_labels.font.name = FONT; va.tick_labels.number_format = '0"%"'; va.tick_labels.number_format_is_linked = False
va.major_gridlines.format.line.color.rgb = LIGHT; va.minimum_scale = -40; va.maximum_scale = 10; va.major_unit = 10
tb(SB, 7100000, 4850000, 4500000, 1100000, [
    [(f'Hotels: {Y_HOT[0] * 100:.0f}% in January to {Y_HOT[8] * 100:.0f}% in September. ', True, NAVY), (f'Water supply: {Y_WAT[0] * 100:.0f}% to {Y_WAT[8] * 100:.1f}%. Commercial total: {Y_COM[0] * 100:.0f}% to {Y_COM[8] * 100:.1f}%.', False, BLACK)],
    [('Storm hit in late October 2025; hotels recover with reopened rooms, water supply with restored pumping. ', False, GREY)]], 9)
tb(SB, 612648, 5900000, 6200000, 150000, 'FY2026 to FY2028 are held at the financials; FY2025 is billed actual. Commercial = RT20 to RT70.', 7.5, False, GREY)
SB.notes_slide.notes_text_frame.text = 'Outlook figures are the held forecast by rate class. The monthly series is billed commercial sales (RT20 to RT70), hotel customers and water supply customers, 2026 against 2025.'


# ================= new slide: key messages =================
assert [n[:8] for x, n in dec6[:3]] == ['National', 'Jamaica ', 'Alcoa Mi'], dec6[:3]
SC, ki3, rr3 = new_from(S5, 'KEY MESSAGES', 'Volume is recovering, but revenue sits with a few large accounts that can leave: the commercial model protects it', '', '')
ki3.text_frame.paragraphs[0].runs[0].text = 'So what and ask'
b3 = ki3.text_frame.paragraphs[-1].runs
b3[0].text = 'Protect the top accounts: '
if len(b3) > 1: b3[1].text = 'risk register in Feb, Tier 1 offers in Mar. Board support requested for flexible contracting.'
for r in b3[2:]: r._r.getparent().remove(r._r)
decl = abs(sum(x for x, n in dec6)) / 1e6
cards = [
    ('1', 'SEPTEMBER', pct(tot[0] / tot[1] - 1), 'commercial sales vs budget',
     [f"Alcoa's move to self-generation explains 6.5 of the {tot[1] - tot[0]:.1f} GWh gap.",
      f"Year to date: {pct(tot[3] / tot[4] - 1)} vs budget, {pct(tot[3] / tot[5] - 1)} vs 2025."]),
    ('2', 'THE BASE', f'{nf_net * 100:.0f}%', 'of key-account revenue retained',
     [f'{logo * 100:.0f}% of accounts and {vol_net * 100:.0f}% of volume stayed.',
      f'Six accounts drive J${decl:,.0f}M of the J${abs(nf26 - nf25) / 1e9:.1f}bn decline: water supply, Jamaica Broilers and Alcoa lead.']),
    ('3', 'AHEAD', pct(ctot[2] / ctot[1] - 1), 'commercial sales in FY2027',
     ['Storm recovery adds 140 GWh; three defections take 71 GWh.', 'J$4.16B across 18 accounts is still at risk.'])]
for i, (num, lab, big, sub, lines) in enumerate(cards):
    x0 = 612648 + i * 3700000
    rect(SC, x0, 1650000, 3560000, 3500000, F_TILE)
    rect(SC, x0, 1650000, 3560000, 70000, CYAN)
    tb(SC, x0 + 160000, 1830000, 3200000, 220000, lab, 10, True, GREY)
    tb(SC, x0 + 160000, 2110000, 3200000, 700000, big, 40, True, NAVY)
    tb(SC, x0 + 160000, 2850000, 3200000, 300000, sub, 11, True, BLUE)
    tb(SC, x0 + 160000, 3350000, 3250000, 1700000, [[(t + ('' if i == len(lines) - 1 else ' '), False, BLACK)] for i, t in enumerate(lines)], 12)
SC.notes_slide.notes_text_frame.text = 'Main message: volume is recovering, but revenue is concentrated in a few large accounts that can leave the grid.'

# ---------------- order: new slides after slide 5 and after slide 6 ----------------
ids = prs.slides._sldIdLst
lst = list(ids)
sa, sb, sc0 = lst[-3], lst[-2], lst[-1]
for e in (sa, sb, sc0): ids.remove(e)
lst = list(ids)
ids.insert(5, sa)            # after the sales slide (position 6)
ids.insert(7, sb)            # after growth and tourism (position 8)
ids.insert(1, sc0)   # key messages after the dashboard
cur = list(ids)
for e in cur: ids.remove(e)
for k in (2, 1, 6, 7, 5, 8, 9, 10, 3, 4): ids.append(cur[k - 1])   # key messages, dashboard, sales, revenue, defection, growth, outlook, model, appendix, service
prs.save(OUT); print('saved', OUT, len(prs.slides))
print('logo', round(logo, 4), 'vol net/gross', round(vol_net, 4), round(vol_gross, 4), 'nf net/gross', round(nf_net, 4), round(nf_gross, 4))
print('YTD commercial vs budget', round(tot[3] / tot[4] - 1, 4), 'vs 2025', round(tot[3] / tot[5] - 1, 4), 'acc', [round(v, 3) for v in acc_tot])
