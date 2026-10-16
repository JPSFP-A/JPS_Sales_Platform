# -*- coding: utf-8 -*-
# Forecast accuracy deck (3 slides) on the JPS Corporate Performance Report template.
# Forecast = budget (account level). Accuracy = 1 - |actual - forecast| / actual. Actuals: Sales Platform (customer billed = 1).
import copy, shutil
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK

SRC = r'C:\Users\jwilson\Downloads\September 2026 Corporate Performance Report.pptx'
OUT = r'C:\Projects\Sales_Platform\JPS_Forecast_Accuracy_Aug_Sep_2026.pptx'
NAVY, BLUE, CYAN, GREY, LIGHT = RGBColor(0x06, 0x25, 0x52), RGBColor(0x0E, 0x4E, 0x95), RGBColor(0x00, 0x9F, 0xDA), RGBColor(0x66, 0x78, 0x8A), RGBColor(0xD9, 0xE1, 0xF2)
RED, GREEN, WHITE, BLACK = RGBColor(0xC0, 0x00, 0x00), RGBColor(0x4E, 0x8B, 0x66), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x1F, 0x1F, 0x1F)
F_GOOD, F_MID, F_BAD, F_ALT = RGBColor(0xDD, 0xEF, 0xDF), RGBColor(0xFF, 0xF0, 0xC8), RGBColor(0xF6, 0xD3, 0xD6), RGBColor(0xF3, 0xF6, 0xFB)
FONT = 'Arial'

# ---------------- data (GWh) ----------------
PER = ('AUG', 'SEP', 'YTD')
CLASS = {   # actual, budget by period
 'RT10': {'AUG': (117.671, 114.154), 'SEP': (114.566, 104.830), 'YTD': (882.601, 853.626)},
 'RT20': {'AUG': (60.610, 58.922), 'SEP': (60.618, 57.649), 'YTD': (475.762, 461.306)},
 'RT40': {'AUG': (71.883, 72.325), 'SEP': (67.313, 73.334), 'YTD': (569.767, 584.913)},
 'RT50': {'AUG': (30.677, 35.090), 'SEP': (30.253, 35.268), 'YTD': (259.549, 274.111)},
 'RT60-ST': {'AUG': (3.413, 3.375), 'SEP': (3.413, 3.312), 'YTD': (27.490, 27.057)},
 'RT70': {'AUG': (19.708, 23.364), 'SEP': (18.643, 22.866), 'YTD': (213.868, 192.712)},
}
ACCT_ABS = {   # sum of absolute account-level errors (accounts with an account-level budget)
 'RT40': {'AUG': 21.609, 'SEP': 22.938, 'YTD': 106.260}, 'RT50': {'AUG': 11.021, 'SEP': 10.841, 'YTD': 60.604},
 'RT60-ST': {'AUG': 0.083, 'SEP': 0.138, 'YTD': 0.726}, 'RT70': {'AUG': 13.486, 'SEP': 13.490, 'YTD': 45.296},
}
IND = {   # industry: per -> (actual, budget, abs account-level error)
 'Water Supply & Irrigation': {'AUG': (19.209, 20.811, 8.338), 'SEP': (18.623, 21.042, 8.774), 'YTD': (158.029, 166.537, 41.227)},
 'Hotels & Motels': {'AUG': (18.232, 20.933, 8.299), 'SEP': (17.231, 20.944, 8.686), 'YTD': (145.540, 156.689, 42.421)},
 'Cement Manufacturing': {'AUG': (8.561, 8.684, 0.123), 'SEP': (9.150, 8.728, 0.444), 'YTD': (82.527, 75.219, 7.373)},
 'Other Metal Ore Mining (incl. Jamalco)': {'AUG': (3.340, 10.446, 7.269), 'SEP': (3.038, 10.258, 7.412), 'YTD': (80.756, 83.915, 5.927)},
 'Telecommunications Carriers': {'AUG': (7.247, 8.238, 2.022), 'SEP': (6.972, 8.189, 2.203), 'YTD': (60.966, 65.744, 10.902)},
 'Executive Offices': {'AUG': (5.631, 5.865, 1.910), 'SEP': (5.931, 5.832, 1.397), 'YTD': (47.379, 46.608, 7.068)},
 'Full-Service Restaurants': {'AUG': (5.380, 5.244, 1.077), 'SEP': (5.138, 5.295, 0.952), 'YTD': (42.369, 42.372, 4.471)},
 'Banking, Insurance & Investments': {'AUG': (4.905, 4.717, 0.726), 'SEP': (4.801, 4.750, 0.744), 'YTD': (38.518, 38.031, 3.496)},
 'Food Manufacturing': {'AUG': (3.879, 3.586, 1.038), 'SEP': (4.407, 3.592, 1.511), 'YTD': (34.282, 29.939, 7.544)},
 'Supermarkets & Grocery Retail': {'AUG': (4.087, 4.007, 0.631), 'SEP': (3.994, 4.050, 0.661), 'YTD': (33.665, 32.315, 4.152)},
 'General Hospitals': {'AUG': (4.424, 3.684, 1.057), 'SEP': (4.433, 3.688, 1.047), 'YTD': (33.461, 29.827, 5.117)},
 'Poultry Hatcheries': {'AUG': (5.351, 4.061, 1.958), 'SEP': (1.993, 4.105, 2.197), 'YTD': (30.413, 33.056, 3.145)},
 'Distilleries': {'AUG': (3.444, 2.616, 1.257), 'SEP': (3.462, 2.588, 1.213), 'YTD': (28.163, 22.897, 7.082)},
}
MAPPED = {'AUG': (137.956, 146.232, 49.556), 'SEP': (132.096, 146.547, 50.401), 'YTD': (1172.644, 1173.177, 222.681)}   # all customers with an industry code

acc = lambda a, f: max(0.0, 1 - abs(a - f) / a) if a else 0.0
pct = lambda x: f'{x * 100:.0f}%'
pct1 = lambda x: f'{x * 100:.1f}%'
sgn = lambda x, d=1: ('+' if x > 0 else ('\u2212' if x < 0 else '')) + f'{abs(x):,.{d}f}'

# ---------------- helpers ----------------
prs = Presentation(SRC)
layout = [l for l in prs.slide_layouts if l.name.strip() == 'Content Slide 1'][0]
# drop every existing slide (keeps masters, layouts, theme)
sld = prs.slides._sldIdLst
for sldId in list(sld):
    prs.part.drop_rel(sldId.rId); sld.remove(sldId)

def tb(slide, x, y, w, h, text='', size=11, bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, italic=False):
    s = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)); tf = s.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0); tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        runs = ln if isinstance(ln, list) else [(ln, bold, color)]
        for t, b, c in runs:
            r = p.add_run(); r.text = t; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = c; r.font.italic = italic
    return s
def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s
def new_slide(n, title, headline, sub):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    tb(s, 711812, 493325, 11000000, 523220, title, 28, True, NAVY)
    tb(s, 815243, 1079336, 11200000, 330000, headline, 16, True, BLUE)
    tb(s, 815243, 1500000, 11200000, 300000, sub, 11, False, GREY)
    tb(s, 469231, 6411345, 6621379, 266676, 'Jamaica Public Service | Forecast Accuracy \u2013 August and September 2026', 9, True, NAVY)
    tb(s, 186473, 6413496, 287140, 265842, str(n), 9, True, NAVY)
    return s
def cell_set(c, text, size=9, bold=False, color=BLACK, align=PP_ALIGN.CENTER, fill=None):
    c.text = ''; tf = c.text_frame; p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    c.margin_left = c.margin_right = Emu(54000); c.margin_top = c.margin_bottom = Emu(18000); c.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is None: c.fill.background()
    else: c.fill.solid(); c.fill.fore_color.rgb = fill
def acc_fill(v): return None if v is None else (F_GOOD if v >= 0.95 else (F_MID if v >= 0.85 else F_BAD))
def make_table(slide, x, y, w, col_w, rows, row_h, header_h=None):
    nr, nc = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nr, nc, Emu(x), Emu(y), Emu(w), Emu(row_h * nr)); t = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for a in ('bandRow', 'firstRow'): tblPr.set(a, '0')
    sid = tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
    if sid is not None: sid.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'     # No Style, Table Grid -> we colour explicitly
    tot = sum(col_w)
    for j, cw in enumerate(col_w): t.columns[j].width = Emu(int(w * cw / tot))
    for i in range(nr): t.rows[i].height = Emu(header_h if (i == 0 and header_h) else row_h)
    return t

# ============================ slide 1: rate class ============================
s1 = new_slide(1, 'Forecast Accuracy by Rate Class',
               'Total sales were forecast within 1.5%, but RT50 and RT70 missed by 14\u201323% in August and September',
               'Forecast = budget. Accuracy = 1 \u2212 |actual \u2212 forecast| \u00f7 actual, by month and for January\u2013September. Volume in GWh, customer-billed basis.')
hdr = ['Rate class', 'Sept actual', 'Sept forecast', 'Sept variance', 'Aug', 'Sept', 'YTD', 'Sept', 'YTD']
rows = [hdr]
names = {'RT10': 'RT10 Residential', 'RT20': 'RT20 Small commercial', 'RT40': 'RT40 Commercial', 'RT50': 'RT50 Large commercial', 'RT60-ST': 'RT60 Street lighting', 'RT70': 'RT70 Industrial'}
order = ['RT10', 'RT20', 'RT40', 'RT50', 'RT60-ST', 'RT70']
body = []
for c in order:
    a, f = CLASS[c]['SEP']; d = CLASS[c]
    ac = [acc(*d[p]) for p in PER]
    aa = [None if c not in ACCT_ABS else max(0, 1 - ACCT_ABS[c][p] / d[p][0]) for p in ('SEP', 'YTD')]
    body.append((names[c], a, f, a - f, ac, aa))
tot = {p: (sum(CLASS[c][p][0] for c in order), sum(CLASS[c][p][1] for c in order)) for p in PER}
A_SEP, F_SEP = tot['SEP']
acct_com_abs = {p: sum(ACCT_ABS[c][p] for c in ACCT_ABS) for p in PER}
acct_com_a = {p: sum(CLASS[c][p][0] for c in ACCT_ABS) for p in PER}
# table with a two-row header (group labels) built as 2 header rows
gx, gy = 274320, 2000000
col_w = [2.3, 1.0, 1.05, 1.05, 0.8, 0.8, 0.8, 0.9, 0.9]
tw = 7650000
t = make_table(s1, gx, gy, tw, col_w, [[''] * 9] * (len(body) + 3), 340000, 300000)
group = [('', 0, 0), ('September volume (GWh)', 1, 3), ('Net accuracy', 4, 6), ('Account-level accuracy', 7, 8)]
for text, c0, c1 in group:
    if text == '':
        for j in (0,): cell_set(t.cell(0, j), '', fill=NAVY)
        continue
    t.cell(0, c0).merge(t.cell(0, c1)); cell_set(t.cell(0, c0), text, 9, True, WHITE, PP_ALIGN.CENTER, NAVY)
for j, h in enumerate(['Rate class', 'Actual', 'Forecast', 'Variance', 'Aug', 'Sept', 'YTD', 'Sept', 'YTD']):
    cell_set(t.cell(1, j), h, 9, True, WHITE, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, BLUE)
for i, (nm, a, f, d_, ac, aa) in enumerate(body, 2):
    base = F_ALT if i % 2 == 1 else None
    cell_set(t.cell(i, 0), nm, 9, True, NAVY, PP_ALIGN.LEFT, base)
    cell_set(t.cell(i, 1), f'{a:,.1f}', 9, False, BLACK, fill=base); cell_set(t.cell(i, 2), f'{f:,.1f}', 9, False, BLACK, fill=base)
    cell_set(t.cell(i, 3), sgn(d_), 9, True, RED if d_ < 0 else BLUE, fill=base)
    for j, v in enumerate(ac): cell_set(t.cell(i, 4 + j), pct(v), 9, True, BLACK, fill=acc_fill(v))
    for j, v in enumerate(aa): cell_set(t.cell(i, 7 + j), 'n/a' if v is None else pct(v), 9, v is not None, GREY if v is None else BLACK, fill=None if v is None else acc_fill(v))
ri = len(body) + 2
cell_set(t.cell(ri, 0), 'Total system', 9, True, WHITE, PP_ALIGN.LEFT, NAVY)
cell_set(t.cell(ri, 1), f'{A_SEP:,.1f}', 9, True, WHITE, fill=NAVY); cell_set(t.cell(ri, 2), f'{F_SEP:,.1f}', 9, True, WHITE, fill=NAVY)
cell_set(t.cell(ri, 3), sgn(A_SEP - F_SEP), 9, True, WHITE, fill=NAVY)
for j, p in enumerate(PER): cell_set(t.cell(ri, 4 + j), pct1(acc(*tot[p])), 9, True, WHITE, fill=NAVY)
for j, p in enumerate(('SEP', 'YTD')): cell_set(t.cell(ri, 7 + j), pct1(1 - acct_com_abs[p] / acct_com_a[p]), 9, True, WHITE, fill=NAVY)
tb(s1, gx, gy + 340000 * (len(body) + 2) + 300000 + 70000, tw, 330000,
   'Account-level accuracy = 1 \u2212 \u03a3|account actual \u2212 account forecast| \u00f7 \u03a3 actual, for RT40\u2013RT70 (accounts with an account-level budget); the total row covers those four classes. '
   'RT10 and RT20 are forecast in consumption buckets, so there is no account-level view.', 8, False, GREY)
# chart: net accuracy by class, Aug / Sept / YTD
cx, cy, cw, ch = 8170000, 2000000, 3800000, 3150000
rect(s1, cx - 80000, cy - 40000, cw + 160000, ch + 80000, None, LIGHT, 0.9)
rect(s1, cx - 80000, cy - 40000, cw + 160000, 45719, BLUE)
tb(s1, cx, cy + 40000, cw, 200000, 'NET FORECAST ACCURACY BY RATE CLASS', 10, True, GREY)
rorder = list(reversed(order))     # horizontal bars plot bottom-up, so reverse to read RT10 first
cd = CategoryChartData(); cd.categories = [c.replace('-ST', '') for c in rorder]
for lab, per in (('YTD', 'YTD'), ('Sept', 'SEP'), ('Aug', 'AUG')): cd.add_series(lab, [acc(*CLASS[c][per]) for c in rorder])
gfc = s1.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Emu(cx), Emu(cy + 260000), Emu(cw), Emu(ch - 280000), cd); chart = gfc.chart
chart.has_title = False; chart.has_legend = True; chart.legend.include_in_layout = False
from pptx.enum.chart import XL_LEGEND_POSITION
chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.font.size = Pt(9); chart.legend.font.name = FONT
plot = chart.plots[0]; plot.gap_width = 45; plot.overlap = -5
for ser, col in zip(plot.series, (NAVY, BLUE, CYAN)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col; ser.invert_if_negative = False
plot.has_data_labels = True; dl = plot.data_labels; dl.number_format = '0%'; dl.number_format_is_linked = False
dl.font.size = Pt(8); dl.font.name = FONT; dl.position = XL_LABEL_POSITION.OUTSIDE_END
ca = chart.category_axis; ca.tick_labels.font.size = Pt(9); ca.tick_labels.font.name = FONT; ca.format.line.color.rgb = LIGHT; ca.has_major_gridlines = False
va = chart.value_axis; va.has_major_gridlines = False; va.minimum_scale = 0; va.maximum_scale = 1.15; va.tick_label_position = XL_TICK_LABEL_POSITION.NONE; va.major_tick_mark = XL_TICK_MARK.NONE; va.format.line.fill.background()
# bottom callout
by = 5560000
rect(s1, gx, by, 11700000 - gx + 274320, 640000, LIGHT)
rect(s1, gx, by, 54000, 640000, BLUE)
tb(s1, gx + 160000, by + 70000, 11200000, 520000, [[('What it means: ', True, NAVY), (
   f'system volume is forecast well (September {pct1(acc(*tot["SEP"]))}), partly because the errors offset: RT10 and RT20 ran above forecast while RT40, RT50 and RT70 ran below. '
   'Much of the shortfall sits in a small number of large accounts. YTD accuracy is flattered: the platform budget for January and February equals actuals.', False, BLACK)]], 10.5)
s1.notes_slide.notes_text_frame.text = ('Forecast = budget (jps_budget_class_month / jps_budget). Actuals: Sales Platform, customer billed = 1, September restated. '
   'Accuracy = 1 - |A-F|/A. YTD = Jan-Sep; the platform budget for Jan and Feb is set equal to actuals, so YTD accuracy is overstated by roughly 1 point.')

# ============================ slide 2: industry ============================
tot_com = {p: (sum(CLASS[c][p][0] for c in ('RT20', 'RT40', 'RT50', 'RT60-ST', 'RT70')), sum(CLASS[c][p][1] for c in ('RT20', 'RT40', 'RT50', 'RT60-ST', 'RT70'))) for p in PER}
mapped_acc_sep = 1 - MAPPED['SEP'][2] / MAPPED['SEP'][0]
mapped_net_sep = acc(MAPPED['SEP'][0], MAPPED['SEP'][1])
s2 = new_slide(2, 'Forecast Accuracy by Industry',
               f'Industry totals look reasonable, but account-level accuracy in September was {pct(mapped_acc_sep)}',
               'Commercial customers with an industry code (RT20\u2013RT70). Forecast = account-level budget. Largest 13 industries by January\u2013September volume.')
hdr2 = ['Industry', 'Sept actual (GWh)', 'Sept variance (GWh)', 'Aug', 'Sept', 'YTD', 'Sept', 'YTD']
t2x, t2y, t2w = 274320, 1950000, 8050000
rows2 = []
for nm, d in IND.items():
    a, f, ab = d['SEP']
    rows2.append((nm, a, a - f, [acc(d[p][0], d[p][1]) for p in PER], [max(0, 1 - d[p][2] / d[p][0]) for p in ('SEP', 'YTD')]))
top_a = {p: sum(IND[n][p][0] for n in IND) for p in PER}; top_f = {p: sum(IND[n][p][1] for n in IND) for p in PER}; top_ab = {p: sum(IND[n][p][2] for n in IND) for p in PER}
oth = {p: (MAPPED[p][0] - top_a[p], MAPPED[p][1] - top_f[p], MAPPED[p][2] - top_ab[p]) for p in PER}
na = {p: (tot_com[p][0] - MAPPED[p][0], tot_com[p][1] - MAPPED[p][1]) for p in PER}
rows2.append(('All other industries', oth['SEP'][0], oth['SEP'][0] - oth['SEP'][1], [acc(oth[p][0], oth[p][1]) for p in PER], [max(0, 1 - oth[p][2] / oth[p][0]) for p in ('SEP', 'YTD')]))
rows2.append(('No industry code (incl. small RT20)', na['SEP'][0], na['SEP'][0] - na['SEP'][1], [acc(na[p][0], na[p][1]) for p in PER], [None, None]))
nrow = len(rows2) + 3
t2 = make_table(s2, t2x, t2y, t2w, [3.2, 0.95, 0.95, 0.65, 0.65, 0.65, 0.8, 0.8], [[''] * 8] * nrow, 222000, 240000)
for text, c0, c1 in (('', 0, 2), ('Net accuracy', 3, 5), ('Account-level accuracy', 6, 7)):
    t2.cell(0, c0).merge(t2.cell(0, c1)); cell_set(t2.cell(0, c0), text, 9, True, WHITE, PP_ALIGN.CENTER, NAVY)
for j, h in enumerate(['Industry', 'Sept actual', 'Sept variance', 'Aug', 'Sept', 'YTD', 'Sept', 'YTD']):
    cell_set(t2.cell(1, j), h, 9, True, WHITE, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, BLUE)
for i, (nm, a, v, ac, aa) in enumerate(rows2, 2):
    base = F_ALT if i % 2 == 1 else None
    cell_set(t2.cell(i, 0), nm, 8.5, True, NAVY, PP_ALIGN.LEFT, base)
    cell_set(t2.cell(i, 1), f'{a:,.1f}', 8.5, False, BLACK, fill=base)
    cell_set(t2.cell(i, 2), sgn(v), 8.5, True, RED if v < 0 else BLUE, fill=base)
    for j, x in enumerate(ac): cell_set(t2.cell(i, 3 + j), pct(x), 8.5, True, BLACK, fill=acc_fill(x))
    for j, x in enumerate(aa): cell_set(t2.cell(i, 6 + j), 'n/a' if x is None else pct(x), 8.5, x is not None, GREY if x is None else BLACK, fill=None if x is None else acc_fill(x))
r_ = nrow - 1
cell_set(t2.cell(r_, 0), 'Commercial total (RT20\u2013RT70)', 8.5, True, WHITE, PP_ALIGN.LEFT, NAVY)
cell_set(t2.cell(r_, 1), f'{tot_com["SEP"][0]:,.1f}', 8.5, True, WHITE, fill=NAVY); cell_set(t2.cell(r_, 2), sgn(tot_com['SEP'][0] - tot_com['SEP'][1]), 8.5, True, WHITE, fill=NAVY)
for j, p in enumerate(PER): cell_set(t2.cell(r_, 3 + j), pct1(acc(*tot_com[p])), 8.5, True, WHITE, fill=NAVY)
for j in (6, 7): cell_set(t2.cell(r_, j), '', 8.5, True, WHITE, fill=NAVY)
# right panel
px, py, pw = 8650000, 1950000, 3420000
rect(s2, px, py, pw, 4180000, None, LIGHT, 0.9); rect(s2, px, py, pw, 45719, CYAN)
tb(s2, px + 130000, py + 110000, pw - 260000, 200000, 'WHAT STANDS OUT', 10, True, GREY)
hot = IND['Hotels & Motels']['SEP']; wat = IND['Water Supply & Irrigation']['SEP']; met = IND['Other Metal Ore Mining (incl. Jamalco)']['SEP']; pou = IND['Poultry Hatcheries']['SEP']
pts = [
 ('1', 'Net hides the account misses', f'Water supply is {pct(acc(*IND["Water Supply & Irrigation"]["YTD"][:2]))} accurate YTD in total, but {pct(1 - IND["Water Supply & Irrigation"]["YTD"][2] / IND["Water Supply & Irrigation"]["YTD"][0])} account by account.'),
 ('2', 'Hotels running below plan', f'{sgn(hot[0] - hot[1])} GWh in September ({pct(acc(hot[0], hot[1]))} accurate); recovery after Hurricane Melissa is the likely cause.'),
 ('3', 'Two accounts, one-off events', 'Jamalco self-generation (\u22126.5 GWh) and a Caribbean Broilers billing credit (\u22121.8 GWh) drive mining and poultry.'),
 ('4', 'Well forecast', 'Restaurants, banking, executive offices and supermarkets are 96\u2013100% accurate YTD.'),
]
yy = py + 380000
for n, h, body_ in pts:
    c = rect(s2, px + 130000, yy, 250000, 250000, BLUE, None, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = n; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    tb(s2, px + 470000, yy - 10000, pw - 600000, 900000, [[(h, True, NAVY)], [(body_, False, BLACK)]], 9.5)
    yy += 930000
leg_y = py + 4180000 - 470000
for k, (lab, col) in enumerate((('95% or better', F_GOOD), ('85–95%', F_MID), ('below 85%', F_BAD))):
    rect(s2, px + 130000 + k * 1050000, leg_y, 150000, 150000, col, LIGHT, 0.5); tb(s2, px + 130000 + k * 1050000 + 200000, leg_y - 10000, 850000, 170000, lab, 8, False, GREY)
tb(s2, px + 130000, leg_y + 210000, pw - 260000, 250000, 'Net = industry total vs forecast; account-level adds errors account by account. 0% = forecast more than double the actual.', 8, False, GREY)
s2.notes_slide.notes_text_frame.text = ('Industry from jps_industries (customer code to industry); 1,002 customers coded. Budget and actuals joined at account level (jps_budget, jps_actuals). '
   'Net accuracy = 1 - |A-F|/A on the industry total; account-level = 1 - sum|a_i - f_i| / sum a_i. Industry totals use customers with an industry code only; the remainder (including bucketed RT20) is shown as "No industry code".')

# ============================ slide 3: implications ============================
com_sep = {c: CLASS[c]['SEP'][0] - CLASS[c]['SEP'][1] for c in ('RT40', 'RT50', 'RT60-ST', 'RT70')}; com_var = sum(com_sep.values())
two = 6.514 + 1.822
com_acct_sep = 1 - acct_com_abs['SEP'] / acct_com_a['SEP']
s3 = new_slide(3, 'What This Means for Customer Experience',
               'Misses sit in a few large accounts and in hospitality; account-level engagement is the fix',
               'September 2026. Forecast = budget; commercial = RT40, RT50, RT60 and RT70.')
tiles = [
 (pct(acc(*tot['SEP'])), 'SYSTEM SALES, SEPTEMBER', f'forecast accuracy, net ({sgn(A_SEP - F_SEP)} GWh vs forecast)'),
 (pct(com_acct_sep), 'LARGE-ACCOUNT FORECAST, SEPTEMBER', 'account-level accuracy, RT40\u2013RT70'),
 (pct(two / abs(com_var)), 'OF THE COMMERCIAL SHORTFALL', f'came from two accounts: Jamalco and Caribbean Broilers ({sgn(com_var)} GWh in total)'),
]
tx, tyy, tw3, th3, gap = 274320, 1950000, 3830000, 1200000, 150000
for k, (big, lab, sub_) in enumerate(tiles):
    x = tx + k * (tw3 + gap)
    rect(s3, x, tyy, tw3, th3, None, LIGHT, 0.9); rect(s3, x, tyy, tw3, 45719, BLUE)
    tb(s3, x + 130000, tyy + 110000, tw3 - 260000, 170000, lab, 9.5, True, GREY)
    tb(s3, x + 130000, tyy + 330000, tw3 - 260000, 420000, big, 30, True, NAVY)
    tb(s3, x + 130000, tyy + 840000, tw3 - 260000, 330000, sub_, 9.5, False, BLACK)
# two columns: drivers and actions
col_y = 3330000; col_h = 2850000; cw3 = 5820000
for k, (head, items) in enumerate((
    ('WHAT IS DRIVING THE MISSES', [
      ('Customers changed what they use', 'Jamalco moved to self-generation (0.8 GWh in September against about 9 GWh a month before), 6.5 GWh below forecast on one account.'),
      ('Hospitality is below plan', f'Hotels were {abs(hot[0] - hot[1]):.1f} GWh below forecast in September; Half Moon Bay and Bush Cay West alone are 1.7 GWh of it.'),
      ('Residential demand is above plan', f'RT10 was {sgn(CLASS["RT10"]["SEP"][0] - CLASS["RT10"]["SEP"][1])} GWh ({(CLASS["RT10"]["SEP"][0] / CLASS["RT10"]["SEP"][1] - 1) * 100:.0f}%) above forecast in September and RT20 {sgn(CLASS["RT20"]["SEP"][0] - CLASS["RT20"]["SEP"][1])} GWh.')]),
    ('WHERE CUSTOMER EXPERIENCE CAN HELP', [
      ('Keep the large-account forecast current', 'Ask key account managers to update the outlook for the 25 largest accounts when a change is known: self-generation, outage, expansion or a billing correction.'),
      ('Engage hotels and water supply early', 'Plan conversations with the largest accounts in these two industries: they are the largest by volume and ran at about 50% account-level accuracy in September.'),
      ('Flag billing credits before they land', 'One billing credit took Caribbean Broilers to 0.1 GWh against 1.9 GWh forecast; share expected credits with planning.')]))):
    x = tx + k * (cw3 + 150000 + 0)
    rect(s3, x, col_y, cw3 + (0 if k == 0 else 0), col_h, None, LIGHT, 0.9); rect(s3, x, col_y, cw3, 45719, CYAN if k else BLUE)
    tb(s3, x + 130000, col_y + 100000, cw3 - 260000, 190000, head, 10, True, GREY)
    yy = col_y + 360000
    for n, (h, b) in enumerate(items, 1):
        c = rect(s3, x + 130000, yy, 240000, 240000, BLUE if k == 0 else CYAN, None, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(n); r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
        tb(s3, x + 480000, yy - 10000, cw3 - 620000, 800000, [[(h, True, NAVY)], [(b, False, BLACK)]], 10.5)
        yy += 810000
s3.notes_slide.notes_text_frame.text = ('Two-account share: Alcoa Mins Of Ja (Jamalco, RT70) -6.5 GWh and Caribbean Broilers (RT40, billing credit) -1.8 GWh against a commercial shortfall of 15.2 GWh in September. '
   'Account-level accuracy covers RT40-RT70. Recommendations are proposals for discussion.')
prs.save(OUT); print('saved', OUT)
print('SEP total acc', acc(*tot['SEP']), 'com acct acc', com_acct_sep, 'two share', two / abs(com_var), 'mapped acc', mapped_acc_sep, 'net', mapped_net_sep)
