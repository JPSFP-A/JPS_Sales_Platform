# -*- coding: utf-8 -*-
# Updates the Executive Summary and Sales Performance slides (slides 1 and 3 of the 4-slide extract; 3 and 7 of the full deck) of the September 2026 Corporate Performance Report
# for September actuals. Sales basis = customer billed = 1 (Platform / Power BI billed line); budget from jps_budget_class_month.
import io, re, shutil, zipfile
from pptx import Presentation
from pptx.chart.data import CategoryChartData
SRC = r'C:\Users\jwilson\Downloads\September 2026 Corporate Performance Report.pptx'
OUT = r'C:\Projects\Sales_Platform\September 2026 Corporate Performance Report_v3.pptx'
TMP = OUT + '.tmp'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
M = '\u2212'   # the deck's minus sign

# ---- September inputs (kWh) ----
CLS = ['RT10', 'RT20', 'RT40', 'RT50', 'RT60', 'RT70']
BUD = dict(zip(CLS, [104829510.00, 57648939.82, 73334099.92, 35267679.99, 3311979.98, 22866400.01]))
ACT = dict(zip(CLS, [114565971.73, 60617570.66, 67312952.00, 30252675.00, 3413193.00, 18642652.00]))
PY = dict(zip(CLS, [110279239.78, 58277527.10, 68941054.00, 33527203.00, 3408535.00, 27281357.00]))
COM = ['RT40', 'RT50', 'RT60', 'RT70']
tot = lambda d, ks=CLS: sum(d[k] for k in ks)
VAR = {k: ACT[k] - BUD[k] for k in CLS}
g = lambda x, d=1: f'{abs(x) / 1e6:,.{d}f}'
s = lambda x, d=1: (M if x < 0 else '+') + g(x, d)
pc = lambda a, b: (b / a - 1) * 100
sp = lambda x: (M if x < 0 else '+') + f'{abs(x):.1f}%'
YTD_ACT, YTD_BUD, YTD_PY = 2429.038e6, 2393.72e6, 2549.957e6    # Jan-Sep, GWh x 1e6 (monthly totals from jps_actuals / budget)

def set_run(run, text):
    run.text = text
    rPr = run._r.get_or_add_rPr()
    for f in rPr.findall(A + 'solidFill'): rPr.remove(f)      # drop the red "to be updated" colour
def one_run(par, text, size=None):
    runs = par.runs; set_run(runs[0], text)
    if size: runs[0].font.size = size
    for r in runs[1:]: r._r.getparent().remove(r._r)

p = Presentation(SRC)
# =================== slide 3 ===================
sh = [x for x in p.slides[0].shapes if x.name == 'Text Placeholder 3'][0]; ps = sh.text_frame.paragraphs
set_run(ps[0].runs[1], f'September billed sales were {g(tot(ACT))} GWh, {g(tot(VAR))} GWh ({sp(pc(tot(BUD), tot(ACT)))}) below budget and {g(tot(ACT) - tot(PY))} GWh ({sp(pc(tot(PY), tot(ACT)))}) below prior year, '
        'as lower RT40, RT50 and RT70 demand, including Jamalco self-generation, outweighed higher Residential and RT20/SME sales.')
one_run(ps[1], f'YTD sales of {YTD_ACT / 1e6:,.1f} GWh remain {g(YTD_ACT - YTD_BUD)} GWh ({sp(pc(YTD_BUD, YTD_ACT))}) above budget but {g(YTD_ACT - YTD_PY)} GWh ({sp(pc(YTD_PY, YTD_ACT))}) below prior year. '
        'Residential (+29.0 GWh), RT20/SME (+14.5 GWh) and RT70 (+21.2 GWh, narrowing as Jamalco moves to self-generation) remain favorable to plan, '
        'while RT40 (' + M + '15.1 GWh) and RT50 (' + M + '14.6 GWh) are below budget, with RT50 affected by lower hotel demand.')
set_run(ps[2].runs[1], ' September non-fuel revenue was US$31.3M, US$1.2M (+3.9%) favorable to budget, supported by a favorable sales mix, '
        'with higher-yield Residential and SME volumes above plan despite total sales 0.8% below budget.')

# =================== slide 7 ===================
sl = p.slides[2]; by = {x.name: x for x in sl.shapes}
def txt(name, text): one_run(by[name].text_frame.paragraphs[0], text)
com_act, com_bud, com_py = tot(ACT, COM), tot(BUD, COM), tot(PY, COM)
txt('TextBox 13', 'Sales Performance \u2013 September 2026')
txt('Text 3', 'September sales below plan as RT40, RT50 and RT70 demand fell; YTD remains above budget.')
txt('Text 5', 'Residential and SME gains partly offset lower commercial sales, with RT40, RT50 and RT70 driving the monthly variance.')
txt('Rectangle 23', f'{g(tot(ACT))} GWh')
txt('Rectangle 25', f'{s(tot(VAR))} GWh vs Budget');            txt('Rectangle: Rounded Corners 27', sp(pc(tot(BUD), tot(ACT))))
txt('Rectangle 29', f'{s(tot(ACT) - tot(PY))} GWh vs September 2025'); txt('Rectangle: Rounded Corners 31', sp(pc(tot(PY), tot(ACT))))
txt('Rectangle 37', f'{g(com_act)} GWh')
txt('Rectangle 39', f'{s(com_act - com_bud)} GWh vs Budget');    txt('Rectangle: Rounded Corners 41', sp(pc(com_bud, com_act)))
txt('Rectangle 43', f'{s(com_act - com_py)} GWh vs September 2025'); txt('Rectangle: Rounded Corners 45', sp(pc(com_py, com_act)))
txt('Rectangle 51', f'{YTD_ACT / 1e6:,.1f} GWh')
txt('Rectangle 53', f'{s(YTD_ACT - YTD_BUD)} GWh vs Budget');   txt('Rectangle: Rounded Corners 55', sp(pc(YTD_BUD, YTD_ACT)))
txt('Rectangle 57', f'{s(YTD_ACT - YTD_PY)} GWh vs Jan-Sep 2025'); txt('Rectangle: Rounded Corners 59', sp(pc(YTD_PY, YTD_ACT)))
txt('Rectangle 63', 'SEPTEMBER SALES BRIDGE'); txt('Rectangle 69', 'SEPTEMBER RATE-CLASS VARIANCE')
pp = by['Rectangle 77'].text_frame.paragraphs
one_run(pp[0], 'SEPTEMBER PERFORMANCE')
one_run(pp[2], f'September sales were {g(tot(VAR))} GWh ({sp(pc(tot(BUD), tot(ACT)))}) below budget and {g(tot(ACT) - tot(PY))} GWh ({sp(pc(tot(PY), tot(ACT)))}) below prior year, '
        'as lower RT40, RT50 and RT70 demand outweighed Residential and SME gains.')
one_run(pp[4], f'Residential and SME were {g(VAR["RT10"] + VAR["RT20"])} GWh above budget combined, while Commercial was {g(com_act - com_bud)} GWh ({sp(pc(com_bud, com_act))}) below plan.')
set_run(pp[6].runs[0], 'Jamalco')
set_run(pp[6].runs[1], f' grid consumption was 0.8 GWh in September, against approximately 9.1 GWh in July, as self-generation came fully online, '
        f'leaving RT70 {g(VAR["RT70"])} GWh below budget and {g(PY["RT70"] - ACT["RT70"])} GWh below prior year.')
one_run(pp[8], 'HIGHLIGHTS')
one_run(pp[10], f'YTD sales remain {g(YTD_ACT - YTD_BUD)} GWh ({sp(pc(YTD_BUD, YTD_ACT))}) above budget but {g(YTD_ACT - YTD_PY)} GWh ({sp(pc(YTD_PY, YTD_ACT))}) below prior year. '
        'Eleven months after Hurricane Melissa, the prior-year gap remains, although September performance is also being shaped by Jamalco\u2019s move to self-generation.', size=pp[2].runs[0].font.size)
one_run(pp[12], 'RT70 remains 21.2 GWh above budget YTD (25.4 GWh at August), but the shift in Jamalco demand is expected to reduce this favorable position over the remainder of the year.')
one_run(pp[14], 'RT40 and RT50 remain below plan YTD by 15.1 GWh and 14.6 GWh, respectively, with RT50 affected by lower hotel demand. '
        'September RT40 includes a billing credit of about 1.6 GWh on one account, which is a timing item.')
# rate-class variance chart (embedded workbook is rewritten by replace_data)
cd = CategoryChartData(); cd.categories = CLS; cd.add_series('Variance', [VAR[k] / 1e6 for k in CLS])
by['Chart'].chart.replace_data(cd)
# slide is now complete: remove the "To be updated" tag
tag = by['TextBox 44']; tag._element.getparent().remove(tag._element)
p.save(TMP)

# =================== waterfall (chartEx) + fallback picture + embedded workbook ===================
res, sme, comv = VAR['RT10'], VAR['RT20'], com_act - com_bud
vals = [tot(BUD), res, sme, comv, tot(ACT)]
labels = [f'{tot(BUD) / 1e6:.2f}', f'{res / 1e6:.2f}', f'{sme / 1e6:.2f}', f'{comv / 1e6:.2f}', f'{tot(ACT) / 1e6:.2f}']
zin = zipfile.ZipFile(TMP); zout = zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED)
import openpyxl
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == 'ppt/charts/chartEx1.xml':
        x = data.decode('utf-8')
        pts = re.findall(r'<cx:pt idx="(\d)">(-?[\d.]+)</cx:pt>', x)          # numeric points
        for idx, v in ((i, v) for i, v in pts):
            x = x.replace(f'<cx:pt idx="{idx}">{v}</cx:pt>', f'<cx:pt idx="{idx}">{vals[int(idx)]!r}</cx:pt>', 1)
        for old, new in zip(['307.23', '3.51', '1.66', '-8.41', '304.00'], labels):
            assert f'<a:t>{old}</a:t>' in x, old
            x = x.replace(f'<a:t>{old}</a:t>', f'<a:t>{new}</a:t>', 1)
        data = x.encode('utf-8')
    elif item.filename == 'ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx':
        wb = openpyxl.load_workbook(io.BytesIO(data)); ws = wb.active
        for i, v in enumerate(vals, 2): ws.cell(i, 2, v)
        b = io.BytesIO(); wb.save(b); data = b.getvalue()
    elif item.filename == 'ppt/media/image5.png':
        fig, ax = plt.subplots(figsize=(8.8, 3.92), dpi=100)
        names = ['2026 Budget', 'Residential', 'SME', 'Commercial', '2026 Actual']; cols = ['#062552', '#009FDA', '#009FDA', '#0E4E95', '#062552']
        base = 0; lo = []
        for i, v in enumerate(vals):
            if i in (0, 4): ax.bar(i, v / 1e6, color=cols[i], width=0.7); top = v / 1e6; base = v / 1e6
            else:
                start = base; base += v / 1e6
                ax.bar(i, abs(v) / 1e6, bottom=min(start, base), color=cols[i], width=0.7)
        for i, l in enumerate(labels): ax.text(i, (vals[i] / 1e6 if i in (0, 4) else 0) + 0.3, l, ha='center', fontsize=10, fontweight='bold')
        ax.set_ylim(280, 313); ax.set_xticks(range(5)); ax.set_xticklabels(names, fontsize=9); ax.set_yticks([])
        for sp_ in ('top', 'right', 'left'): ax.spines[sp_].set_visible(False)
        b = io.BytesIO(); fig.savefig(b, format='png', bbox_inches='tight'); plt.close(fig); data = b.getvalue()
    zout.writestr(item, data)
zout.close(); zin.close()
import os; os.remove(TMP)
print('saved', OUT)
print('bridge', [round(v / 1e6, 2) for v in vals], 'check', round((vals[0] + vals[1] + vals[2] + vals[3]) / 1e6, 3), round(vals[4] / 1e6, 3))
print('YTD act/bud/py', YTD_ACT / 1e6, YTD_BUD / 1e6, YTD_PY / 1e6)
